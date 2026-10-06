"""One versioned read-side source map: page identity, logical slides, content roles, corrections.

Every consumer that points at "a slide" — navigation, transcripts, media, questions, concepts, retrieval — reads the
same map of one source file instead of assuming that a slide is a page. The map keeps these identities apart:

* the source: id, version and the sha256 of the bytes that were inspected;
* the physical PDF page index (1-based; the base is stated in the data) and, only when the file defines one, the
  PDF page label;
* the printed slide number, only when a short isolated number stands at one place in the page margin across the deck,
  advances like a slide number and, where it repeats, repeats under one heading (the evidence is recorded) — a number
  in the body text, or a counter that stays the same over different slides, is never one;
* the PPTX slide id and ordinal, only when a presentation of the same unit exists and a text/title sequence
  alignment of at least three pages supports the correspondence — never "slide N is page N";
* the app's logical slide: one or more adjacent pages with the evidence for the grouping, the index of each page in
  it, and optional bounded regions on a page (a handout page holding several slides);
* a content role per page, classified separately from grouping. A role never removes a page from the inventory.

A group of pages is one of three kinds, and the label of a page says which:

* ``incremental_build`` ("build 2 of 3"): each page shows all that the page before it shows — its lines in the same
  order or its text at the same places, its pictures, its shapes — and more;
* ``duplicate`` ("copy 2 of 2"): the pages draw exactly the same, mark for mark;
* ``animation_steps`` ("step 2 of 4"): one authored slide in several states — equal or growing text while something
  drawn changes, is covered (a fill in the paper's paint is laid that the earlier page does not have), or could not
  be compared, or the deck's own number says so. Never more than medium confidence, and no page of it is said to
  show everything (``cumulative: false``). The name says what the file shows: that two such pages look different on
  screen is not claimed.

Grouping is proposed from several signals (adjacency, heading, ordered text containment, text position, pictures by
their bytes, drawing operations, printed numbers, PDF labels, PPTX alignment). Equal words alone do not make a group:
a rearranged line is not a build, and the same heading and caption over another picture stay two pages. When the
signals are weak or disagree the pages stay separate and the map says ``grouping: uncertain`` with the reasons.
Nothing is collapsed or deleted: every page stays addressable, and ``page=N`` of a source version is always valid.

What the map does not know: the text of a page is the text embedded in the file. Whether a shape or picture covers
it cannot be told without rendering, so a page may list words it does not show (``text_basis``, and
``text_visibility: not_established`` on every logical slide). Two counts per page point at where that matters: the
fills and the text runs painted in the paint of the page's own background. Nothing here renders a page.

A picture or a drawing never lets a page be read as blank or as a bare heading for want of words. What nearly every
page draws counts as the template only when it stands in the margins, covers the page or is a plain panel; what
most pages draw in the body of the page may be one figure the deck walks through, and stays content.

Corrections (group, split, role, printed number, regions, revert) are immutable revisions under the app's own area;
applying them yields a new map revision, keeps the history, and leaves an alias for every logical-slide id and a
retired entry for every region id that changed. A learner's correction is never replaced or undone by an agent's,
an agent never records again what the learner undid, and a record that does not say who made it is not taken for
the learner's. Where corrections in force disagree, the one that came into force last decides. The computed
map is cached there too, keyed by source sha256 + map version + overrides revision, and a cached file is trusted only
while the digest it carries of its own content matches. A map whose inspection was cut short by something of the moment
(a read error, the time bound, a presentation of the unit that could not be opened) is marked ``transient`` and is
never stored: every call inspects again until the file reads through; ``refresh=True`` always inspects again. A
replaced source gets a new map that lists which earlier ids and corrections no longer apply.

Pure read side: nothing here writes into a course folder, renders a page, runs OCR or calls a model. Everything in a
map is inferred from what the file holds; a group is a proposal until someone has looked.

Public API: :func:`build_source_map`, :func:`load_or_build`, :func:`apply_override`, :func:`describe_page`,
:func:`resolve`, :func:`summary` (plus :func:`inspect_pptx` and :func:`role_policy`).

Shape of a map (plain JSON data)::

    contract, detector_revision, map_revision        "source-map/1", heuristics revision, digest of this revision
    source            {source_id, version_id, sha256 (None when the file was not read), filename, bytes, kind, registration, index_sha256_matches}
    page_index_base   1                              page_count        physical pages (None when not inspected)
    inspection        {status: complete | partial | unreadable | not_inspected | unavailable | pptx_only, reason, tool, limits, pages_unknown,
                       transient, attempts, retry: automatic | on_request}
    page_layout       {kind: slides | document | mixed | unknown, landscape_pages, upright_pages}
    pdf_labels        {defined, differ_from_index, repeat, repeat_is_a_signal, distinct}
    printed_numbers   {established, numbers_count_slides, evidence}
    pptx              None | {source, status, reason, transient, slide_count, hidden_slides, truncated, alignment, candidates, slides}
    text_basis        {source: embedded_text, visible_text_established: false, note}
    pages[]           page, status (ok | unknown), reason (unreadable | read_error | time_limit | page_limit), pdf_label, printed_number (+_source),
                      pptx {slide_id, ordinal, evidence} | None, title, text, word_count, content_word_count, text_sha256, content_sha256,
                      images[] {name, width, height, bytes, sha256, digest_basis, box, template}, image_count, pictures_unread,
                      drawing {paint_operations, path_operations, shapes_beyond_template, fills_in_background_paint, text_runs_in_background_paint,
                               marks_compared, visual_sha256},
                      size, regions[] {region_id, order, box, label, source}, retired_regions[] {region_id, box, retired_by_revision, resolves_to},
                      role {role, confidence, uncertain, evidence, source, forced_requirements}, logical_slide_id, build_index, build_count
    logical_slides[]  logical_slide_id, ordinal, kind (single | incremental_build | duplicate | animation_steps | manual_group),
                      grouping (none | inferred | uncertain | override), confidence (high | medium | None), pages, representative_page,
                      cumulative (true | false | None), text_visibility, title, printed_numbers, pptx_slides, role, role_uncertain,
                      evidence[] {pages, relation, kind, signals, confidence, measures, notes}, uncertain_reasons, uncertain_with, repeats_title_of, overrides
    pairs[]           adjacent pages that have something in common: decision (join | uncertain | distinct), relation, kind, confidence, signals, reasons, measures
    aliases[]         {alias, pages, retired_by_revision, resolves_to, primary}
    overrides         {revision, history[] {revision, type, op, author, reason, created_at, supersedes, replaces,
                       status: applied | superseded | reverted | refused | invalid, in_force, reverted_by, restores}, problems}
    other_source_versions[]   ids and corrections kept for other bytes of the same source (not applied)
    cache             {key: {sha256, map_version, overrides_revision}, inputs_digest, base_digest, analysis_digest, content_sha256} (only from load_or_build)
"""
from __future__ import annotations

from array import array
from collections import Counter
from dataclasses import asdict, dataclass, field
from datetime import datetime, timezone
from importlib import metadata
from pathlib import Path
import copy
import hashlib
import json
import math
import re
import time
import unicodedata
import zipfile
from typing import Any, Callable

from pypdf import PdfReader
from pypdf.generic import IndirectObject

from . import contracts as c
from .config import Settings
from .fsutil import ExclusiveConflict, publish_exclusive, replace_atomic
from .ownerio.inventory import disk_version_id

SOURCE_MAP_VERSION = "source-map/1"
OVERRIDE_CONTRACT = "source-map-override/1"
DETECTOR_REVISION = 3  # bump when the grouping, numbering or role heuristics change: cached maps are then rebuilt
PAGE_INDEX_BASE = 1
REGION_BOX_UNITS = "fraction_of_page_from_top_left"
ROLES = ("instructional", "worked_example_or_question", "objectives_or_summary", "structural_title_or_section", "administrative", "blank", "mixed", "unknown")
# incremental_build: every page shows all that the page before it shows, and more.   duplicate: the pages draw exactly the same.
# animation_steps: one authored slide in several states, where a later page changes, moves or covers something of an earlier one.
GROUP_KINDS = ("incremental_build", "duplicate", "animation_steps", "manual_group")
_GROUP_WORD = {"incremental_build": "build", "duplicate": "copy", "animation_steps": "step"}
AUTHORS = ("learner", "agent")
_NO_FORCED_REQUIREMENT = frozenset({"structural_title_or_section", "administrative", "blank"})
_INSTRUCTIONAL = frozenset({"instructional", "worked_example_or_question", "mixed"})
_PASSING_FAILURES: dict[tuple[str, str, str], int] = {}  # inspections in a row, in this process, that ended in a failure of the moment (never stored)

__all__ = ["SOURCE_MAP_VERSION", "OVERRIDE_CONTRACT", "ROLES", "Limits", "SourceMapError", "build_source_map", "load_or_build", "apply_override",
           "describe_page", "resolve", "summary", "inspect_pptx", "role_policy"]


class SourceMapError(ValueError):
    def __init__(self, code: str, message: str):
        super().__init__(f"{code}: {message}")
        self.code = code
        self.message = message


@dataclass(frozen=True)
class Limits:
    """Bounds of one inspection. A page beyond a bound stays in the inventory as ``unknown``; it is never dropped."""
    max_pages: int = 2000
    max_bytes: int = 300 * 1024 * 1024
    max_seconds: float = 180.0
    max_page_text_chars: int = 20_000
    max_images_per_page: int = 64
    max_pptx_bytes: int = 300 * 1024 * 1024
    max_pptx_expanded_bytes: int = 1_200 * 1024 * 1024
    max_pptx_slides: int = 2000
    max_marks_per_page: int = 6000  # painted shapes and text runs kept for comparison; a page with more is compared by its text only


# ---- text helpers -------------------------------------------------------------------------------------------
_WORD = re.compile(r"[^\W_]+")
_TOKEN = re.compile(r"[^\W_]+|[^\w\s]")
_CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff]")
_BULLET = re.compile(r"^\s*(?:[-–—•·▪■□●○*»>]|\d{1,2}[.)])\s+")
_STROKE_OPS = frozenset({b"S", b"s"})
_FILL_OPS = frozenset({b"f", b"F", b"f*"})
_BOTH_OPS = frozenset({b"B", b"B*", b"b", b"b*"})
_PATH = frozenset({b"m", b"l", b"c", b"v", b"y", b"re"})
_TEXT_SHOW = frozenset({b"Tj", b"TJ", b"'", b'"'})
_FILL_COLOUR = frozenset({b"g", b"rg", b"k", b"sc", b"scn"})
_DEVICE_COLOUR = frozenset({b"g", b"rg", b"k", b"G", b"RG", b"K"})
_WHITE = frozenset({"g1.00", "rg1.00 1.00 1.00", "k0.00 0.00 0.00 0.00"})  # the paint of bare paper, as the fill state words it
_STROKE_COLOUR = frozenset({b"G", b"RG", b"K", b"SC", b"SCN"})
_LINE_STATE = frozenset({b"w", b"d", b"J", b"j", b"M"})
_TEXT_STATE = frozenset({b"Tr", b"Tc", b"Tw", b"Tz", b"Ts"})
_TRAILING_NUMBER = re.compile(r"(?:^|[^\w.,:/])(\d{1,4})(?:\s*(?:/|of)\s*\d{1,4})?\s*$")
# A text run that is nothing but a number, as a slide or page number is printed: "12", "12 / 40", "Slide 12", "- 12 -".
_NUMBER_RUN = re.compile(r"^[-–—|\s]*(?:(?:slide|page|p\.?|#)\s*)?(\d{1,4})(?:\s*(?:/|of)\s*\d{1,4})?[-–—|\s]*$", re.I)
_MARGIN_TOP, _MARGIN_BOTTOM = 0.12, 0.85  # share of the page height, from the top, inside which a printed number is looked for
_LOGICAL_ID = re.compile(r"^ls-([0-9a-f]{10})-p(\d+)(?:-(\d+))?$")
_PAGE_REF = re.compile(r"^(?:(?:page|p)\s*=\s*)?(\d+)(?:\s*(?:#|&region=|#region=)\s*(r\d+))?$")
_SAFE_ID = re.compile(r"^[A-Za-z0-9][A-Za-z0-9._-]{0,100}$")
_VERSION_DIR = re.compile(r"^[0-9a-f]{16}$")
_MAP_FILE = re.compile(r"^map-o(\d{4})(?:-c[0-9a-f]{8})?\.json$")

_DENSE_SLIDE, _DENSE_UPRIGHT = 400, 150  # words beyond which a page is running text, on a landscape / an upright page
# Signs that are only written in a formula, a relation or a reaction.
_FORMULA_MARKS = frozenset("=≈≠≤≥≪≫∑∏∫√∂∇±×÷∝^∞→←↔⇌⇒⇔⟶⟷∆∈∉⊂⊆∪∩≡∼−")
# Operators that are also punctuation count only where they stand between operands: "A + B", "c - b", "v / (k [A])", "p < 0.05", "A -> B".
_RE_FORMULA = re.compile(
    r"->|<-|=>|>>|<<|>=|<="  # arrows and comparisons spelled with two characters
    r"|(?<![<>\w])[<>](?![<>\w])"  # a lone < or > between other things
    r"|[\w\])]\s*[+*]\s*[\w(\[]"  # x + y, x * y
    r"|(?<![A-Za-z])[A-Za-z)\]]\s+-\s+(?:[A-Za-z](?![A-Za-z])|\d|[(\[])|(?<![A-Za-z\d])\d+(?:\.\d+)?\s+-\s+[A-Za-z](?![A-Za-z])"  # x - y, with a one-letter operand
    r"|(?<![A-Za-z])[A-Za-z]\s*/\s*[A-Za-z](?![A-Za-z])|[\w)\]]\s*/\s*[(\[]|[)\]]\s*/\s*\w"  # v / k, v / (k), (a) / b
    r"|(?<![A-Za-z])[A-Za-z][(\[][^()\[\]]{0,40}[)\]]"  # a one-letter name applied to something, with no space between: f(x), P(A | B), E[X], N(0, 1)
    r"|(?<![A-Za-z])(?:Var|Cov|Corr|Pr|exp|log|ln|sin|cos|tan|max|min|det|lim|sup|inf|arg ?max|arg ?min)\s?[(\[]|(?<![A-Za-z])arg ?(?:max|min)(?![A-Za-z])"  # named operators
    r"|(?<![A-Za-z])d[A-Za-z]{1,2}\s*/\s*d[A-Za-z]{1,2}(?![A-Za-z])"  # a derivative: dN/dt
    r"|[A-Za-z]_[A-Za-z0-9]|\S\s~\s\S")  # a subscript written with an underscore, "is distributed as"
# Abbreviations and addresses that only look like a ratio or an equation.
_RE_NOT_FORMULA = re.compile(r"(?<![A-Za-z])(?:q\s*/\s*a|i\s*/\s*o|n\s*/\s*a|a\s*/\s*c|w\s*/\s*o|c\s*/\s*o|a\s*/\s*v|b\s*/\s*w)(?![A-Za-z])|\band\s*/\s*or\b|https?://\S+|www\.\S+|\S+@\S+", re.I)
# "Example: ...", "Question 2", "Exercise 3" — but not "Examples of ...", "Exercise and oxygen uptake", "Problems with ..." or a closing "Questions?".
_RE_QUESTION = re.compile(r"^(?:worked\s+)?(?:examples?\b(?!\s+of\b)|exercises?\b(?=\s*(?:$|\d|[:.#(\-–—]|for you\b))|practice (?:problems?|questions?)\b|"
                          r"problems?\b(?!\s+(?:with|of|in|sets?)\b)|"
                          r"questions?\b(?!\s*\?\s*$)|quiz\b|clicker\b|poll\b|try it\b|your turn\b|think[- ]pair\b|in[- ]class (?:question|exercise|activity|problem)s?\b|"
                          r"check your understanding\b|concept(?:ual)? (?:check|question)s?\b|warm[- ]up\b)", re.I)
_RE_ADMIN_EXACT = re.compile(r"^(?:problem sets?|homeworks?|assignments?|exams?|course staff|staff|contact|textbooks?|prerequisites?|schedule|reminders?|policies|grading)\s*:?$", re.I)
_RE_OPTION = re.compile(r"^\(?[A-Ea-e][).:]\s+\S")
# "grading" alone is also a topic ("Tumor grading and staging"): it names administration only as a course phrase or as the whole heading.
_RE_ADMIN = re.compile(r"\b(?:office hours?|syllabus|grading (?:policy|policies|scheme|breakdown|criteria|scale|rubric)|(?:course|class|homework|exam|assignment) grading|"
                       r"logistics|administrivia|housekeeping|announcements?|course (?:info|information|policies|policy|schedule|staff|website)|"
                       r"important dates|due dates?|deadlines?|late policy|attendance policy|academic (?:integrity|honesty)|exam (?:dates?|schedule|logistics))\b", re.I)
_RE_ADMIN_BODY = re.compile(r"\b(?:office hours?|due (?:on|by|date)|is due|deadline|gradescope|blackboard|piazza|syllabus|late policy|submit(?:ted)? (?:on|by|via))\b", re.I)
_RE_OBJECTIVES = re.compile(r"^(?:(?:learning|lecture|course|today'?s|class) )?(?:objectives?|outcomes?|goals?|outline|agenda|overview|road ?map|plan|contents|topics)\b|"
                            r"^(?:summary|recap|review|conclusions?|key (?:points|takeaways|ideas|concepts)|take[- ]?(?:home|aways?)|what we (?:learned|covered)|in summary|wrap[- ]up)\b", re.I)
# "Review of probability" teaches probability; "Review of last lecture" looks back at the course.
_RE_OF_THE_COURSE = re.compile(r"\s+of\s+(?:(?:this|the|today'?s|our|last|previous|next)\s+)*(?:lectures?|class|course|session|week|unit|module|day|talk|chapter|part|section|term|semester|time)\b", re.I)
_RE_CLOSING = re.compile(r"^(?:(?:any )?questions\s*\??|q\s*[&/]\s*a|thank you\s*!?|thanks\s*!?|the end|end)$", re.I)
_RE_REFERENCES = re.compile(r"^(?:references|bibliography|further reading|acknowledge?ments?|credits)\b", re.I)
_RE_SECTION = re.compile(r"^(?:part|section|chapter|unit|module|lecture|week|topic)\s+(?:\d+|[ivxlc]+)\b", re.I)
_RE_COVER = re.compile(r"\b(?:lecture|week|chapter|unit|module|session)\s+\d+\b|\b(?:19|20)\d{2}\b|\b(?:university|department|professor|prof\.|instructor|fall|spring|semester)\b", re.I)
# What names a course on a first page that also shows a picture or a drawing. A season, a year or "unit 1" alone also occur in the
# heading of a content page ("Free fall", "The 1953 double helix"): there only a lecture number, a term or an institution counts.
_RE_COURSE_HEADER = re.compile(r"\b(?:lecture|week|chapter|module|session|class)\s+\d+\b|\b(?:fall|spring|summer|winter|autumn)\s+(?:term\s+)?(?:19|20)\d{2}\b|"
                               r"\b(?:university|department|professor|prof\.|instructor|semester)\b", re.I)


def _clean(text: str) -> str:
    return _CONTROL.sub(" ", text or "")


def _norm(text: str) -> str:
    return " ".join(unicodedata.normalize("NFKC", text or "").casefold().split())


def _words(text: str) -> list[str]:
    return _WORD.findall(_norm(text))


def _digest(value: Any) -> str:
    return c.sha256_hex(c.dump_json(value))


def _plain(value: Any) -> Any:
    """The value as plain JSON data (lists, dicts, numbers, strings): what a cache read gives back."""
    return json.loads(c.dump_json(value))


_FILE_HASHES: dict[tuple[str, int, int], str] = {}


def _file_sha256(path: Path) -> str:
    stat = path.stat()
    key = (str(path), stat.st_size, stat.st_mtime_ns)
    if key not in _FILE_HASHES:
        digest = hashlib.sha256()
        with open(path, "rb") as handle:
            for chunk in iter(lambda: handle.read(1024 * 1024), b""):
                digest.update(chunk)
        _FILE_HASHES[key] = digest.hexdigest()
    return _FILE_HASHES[key]


def role_policy(role: str) -> dict[str, Any]:
    """What the study flow may expect of a page with this role. Structural, administrative and blank pages carry no
    forced teach-back, generated practice, concept requirement or missing-match warning; they offer the way onward."""
    forced = role not in _NO_FORCED_REQUIREMENT
    return {"teach_back": forced, "generated_practice": forced, "concept_requirement": forced, "missing_match_warning": forced,
            "offer": [] if forced else ["next_content", "whole_lecture"]}


# ---- PDF inspection (pypdf only) --------------------------------------------------------------------------------
def _deref(value: Any) -> Any:
    return value.get_object() if hasattr(value, "get_object") else value


def _box_fraction(matrix: list[float], corners: list[tuple[float, float]], left: float, bottom: float, width: float, height: float) -> list[float] | None:
    if width <= 0 or height <= 0:
        return None
    a, b, cc, d, e, f = matrix
    xs = [a * x + cc * y + e for x, y in corners]
    ys = [b * x + d * y + f for x, y in corners]

    def clamp(value: float) -> float:
        return round(min(1.0, max(0.0, value)), 4)

    return [clamp((min(xs) - left) / width), clamp(1 - (max(ys) - bottom) / height), clamp((max(xs) - left) / width), clamp(1 - (min(ys) - bottom) / height)]


_IDENTITY = (1.0, 0.0, 0.0, 1.0, 0.0, 0.0)
_NO_BOX = (math.nan, math.nan, math.nan, math.nan)  # a mark that has no place of its own on the page
_GROUND_PICTURE = "picture"  # stands for the background's paint once a picture covers the page: no paint equals it
_PAGE_COVER = 0.9  # share of the page from which a fill or a picture is the page's background
_PANEL = 0.2  # share of the page from which a plain closed box is a panel or a frame, not part of a figure
_FIGURE = 0.02  # share of the page that what most pages draw in the body must span to be read as a figure rather than as the template
_BODY = (0.15, 0.22, 0.85, 0.85)  # the body of a page (left, top, right, bottom as shares of the page): inside the margins, below the heading


def _mul(first: Any, second: Any) -> tuple[float, float, float, float, float, float]:
    """The transformation that applies ``first`` and then ``second`` (PDF matrices, six numbers each)."""
    a1, b1, c1, d1, e1, f1 = (float(v) for v in first)
    a2, b2, c2, d2, e2, f2 = (float(v) for v in second)
    return (a1 * a2 + b1 * c2, a1 * b2 + b1 * d2, c1 * a2 + d1 * c2, c1 * b2 + d1 * d2, e1 * a2 + f1 * c2 + e2, e1 * b2 + f1 * d2 + f2)


def _area(box: Any) -> float:
    return max(0.0, box[2] - box[0]) * max(0.0, box[3] - box[1])


def _in_body(box: Any) -> bool:
    """Whether a box (fractions of the page from the top left) is centred in the body of the page and is not the page's background."""
    if box is None or any(math.isnan(v) for v in box) or _area(box) >= _PAGE_COVER:
        return False
    x, y = (box[0] + box[2]) / 2, (box[1] + box[3]) / 2
    return _BODY[0] <= x <= _BODY[2] and _BODY[1] <= y <= _BODY[3]


def _stream_bytes(obj: Any) -> bytes | None:
    """The stored bytes of a stream object, or ``None`` when the parser does not hand them out."""
    raw = getattr(obj, "_data", None)
    if isinstance(raw, (bytes, bytearray)):
        return bytes(raw)
    read = getattr(obj, "get_data", None)  # the documented way: it decodes the stream, so it is the second choice
    if callable(read):
        try:
            data = read()
        except Exception:
            return None
        return bytes(data) if isinstance(data, (bytes, bytearray)) else None
    return None


def _raw(table: Any, name: Any) -> Any:
    """The entry of a resource table as it is stored (a reference, when it is one), so that it can be remembered by its object number."""
    getter = getattr(table, "raw_get", None)
    return getter(name) if callable(getter) else table[name]


def _object_key(value: Any) -> tuple[int, int] | None:
    ref = value if isinstance(value, IndirectObject) else getattr(value, "indirect_reference", None)
    return (ref.idnum, ref.generation) if isinstance(ref, IndirectObject) else None


_PICTURE_LOOK = ("/Width", "/Height", "/BitsPerComponent", "/ColorSpace", "/Decode", "/DecodeParms", "/Filter", "/ImageMask", "/Mask", "/SMask")


def _image_record(name: str, ref: Any, box: list[float] | None, memo: dict | None = None) -> dict:
    """One picture. ``digest_basis`` says what the digest stands on: only ``stream_bytes`` and ``inline_bytes``
    identify the picture; ``name_and_size`` is a placeholder that is never taken as evidence that two pictures are equal.

    The digest covers the stored bytes together with what decides how they are shown — size, colour space (with its
    palette), decode array, filter, masks — so the same bytes shown inverted, through another palette or under
    another mask are another picture. When one of those cannot be read the picture is not identified at all.
    """
    key = _object_key(ref)
    obj = _deref(ref)
    known = memo.get(("picture", key)) if memo is not None and key is not None else None
    if known is None:
        raw = _stream_bytes(obj)
        try:
            look = "|".join(f"{entry}:{_canon(obj[entry])}" for entry in _PICTURE_LOOK if entry in obj) if raw is not None else None
        except Exception:  # a palette, a mask or a parameter that cannot be read
            look = None
        if raw is not None and look is not None:
            digest = hashlib.sha256(raw)
            digest.update(b"|" + look.encode("utf-8", "backslashreplace"))
            known = (digest.hexdigest(), len(raw), "stream_bytes")
        else:
            known = (None, None, "name_and_size")
        if memo is not None and key is not None:
            memo[("picture", key)] = known
    digest, size, basis = known
    width, height = obj.get("/Width"), obj.get("/Height")
    if digest is None:
        digest = hashlib.sha256(f"{name}|{width}x{height}".encode("utf-8")).hexdigest()
    return {"name": name.lstrip("/"), "width": int(width) if width is not None else None, "height": int(height) if height is not None else None,
            "bytes": size, "sha256": digest, "digest_basis": basis, "box": box, "template": False}


def _picture_key(image: dict, page: int, index: int) -> str:
    """What a picture is compared by: its bytes, or a key no other picture shares when its bytes are not known."""
    return image["sha256"] if image["digest_basis"] != "name_and_size" else f"unknown:{page}:{index}"


def _form_images(form: Any, box: list[float] | None, depth: int, out: list[dict], limit: int, memo: dict | None = None) -> None:
    if depth > 3 or len(out) >= limit:
        return
    resources = _deref(form.get("/Resources")) or {}
    table = _deref(resources.get("/XObject")) or {}
    for name in table:
        ref = _raw(table, name)
        obj = _deref(ref)
        subtype = obj.get("/Subtype")
        if subtype == "/Image" and len(out) < limit:
            out.append(_image_record(str(name), ref, box, memo))
        elif subtype == "/Form":
            _form_images(obj, box, depth + 1, out, limit, memo)


def _fmt(values: Any) -> str:
    """Operands as text: numbers to two decimals, strings as they are. The same operands always give the same text."""
    out = []
    for value in values:
        if isinstance(value, bool):
            out.append(str(value))
        elif isinstance(value, (int, float)):
            out.append("%.2f" % (round(float(value), 2) + 0.0))
        elif isinstance(value, (bytes, bytearray)):
            out.append(bytes(value).hex())
        elif isinstance(value, (list, tuple)):
            out.append("[" + _fmt(value) + "]")
        else:
            out.append(str(value))
    return " ".join(out)


_CANON_SKIP = frozenset({"/Parent", "/Length", "/StructParents", "/StructParent", "/Metadata", "/PieceInfo", "/LastModified", "/OC", "/Name"})


def _canon(obj: Any, depth: int = 0) -> str:
    """A resource object as text that does not depend on its name, its object number or the run: what it holds."""
    obj = _deref(obj)
    if isinstance(obj, dict):
        if "/BaseFont" in obj:
            return f"font:{obj.get('/BaseFont')}:{obj.get('/Subtype')}"
        if depth >= 3:
            return "{}"
        is_stream = callable(getattr(obj, "get_data", None))
        data = _stream_bytes(obj) if is_stream else None
        if is_stream and data is None:  # never taken for equal to another: the caller gives up comparing the page's drawing
            raise ValueError("the bytes of a stream could not be read")
        head = "" if data is None else "stream:" + hashlib.sha256(data).hexdigest()[:20]
        return head + "{" + ",".join(f"{key}:{_canon(value, depth + 1)}" for key, value in sorted(obj.items(), key=lambda item: str(item[0])) if key not in _CANON_SKIP) + "}"
    if isinstance(obj, (list, tuple)):
        return "[" + ",".join(_canon(value, depth + 1) for value in obj[:64]) + "]"
    if isinstance(obj, bool):
        return str(obj)
    if isinstance(obj, (int, float)):
        return "%.4g" % obj
    if isinstance(obj, (bytes, bytearray)):  # a long string (a palette) by a digest of all of it
        return bytes(obj).hex() if len(obj) <= 64 else "sha:" + hashlib.sha256(bytes(obj)).hexdigest()[:32]
    return str(obj)


def _identity(value: Any, memo: dict) -> str:
    """A short digest of a resource (a graphics state, a font, a shading, a form), remembered per stored object."""
    key = _object_key(value)
    if key is not None and key in memo:
        return memo[key]
    digest = hashlib.sha256(_canon(value).encode("utf-8", "backslashreplace")).hexdigest()[:16]
    if key is not None:
        memo[key] = digest
    return digest


def _title_from(chunks: list[tuple], lines: list[str]) -> str | None:
    """The heading of a page: the run of largest text when the page sets it apart by size (with whatever shares its
    baseline, such as an italic word set a little smaller), else the first line. A page that holds nothing but a few
    lines of one size (a cover, a divider) is all heading: a heading that wraps is then read whole."""
    def lettered(index: int) -> bool:
        return sum(ch.isalpha() for ch in chunks[index][0]) >= 2

    candidates = [i for i in range(len(chunks)) if lettered(i)]
    if candidates:
        top = max(chunks[i][1] for i in candidates)
        first = next(i for i in candidates if chunks[i][1] >= top * 0.92)
        picked, baselines = [first], [chunks[first][2]]
        if top > 0:
            for i in range(first + 1, len(chunks)):
                large = chunks[i][1] >= top * 0.92
                if not large and not any(abs(chunks[i][2] - line) <= 0.35 * top for line in baselines):
                    break
                picked.append(i)
                if large:
                    baselines.append(chunks[i][2])
            start = first
            while start > 0 and abs(chunks[start - 1][2] - chunks[first][2]) <= 0.35 * top:
                start -= 1
                picked.insert(0, start)
        title = " ".join(" ".join(chunks[i][0] for i in picked).split())
        apart = any(i not in picked for i in candidates)  # the page also holds smaller text: the large run is its heading
        few = len(lines) <= 4 and not any(_BULLET.match(line) for line in lines)  # a cover or a divider, not a list
        if len(title) <= 160 and (apart or few):
            return title
    first_line = next((line for line in lines if sum(ch.isalpha() for ch in line) >= 2), None)
    return " ".join(first_line.split())[:200] if first_line else None


class _TimeLimit(Exception):
    """The inspection ran past its time bound in the middle of a page."""


def _unknown_page(number: int, reason: str) -> dict:
    return {"page": number, "status": "unknown", "reason": reason, "text": "", "text_truncated": False, "word_count": 0, "title": None, "images": [], "image_count": 0,
            "images_truncated": False, "paint_operations": None, "path_operations": None, "text_operations": None, "text_in_background_paint": None, "size": None, "origin": None,
            "placements": [], "chunks": [], "vector_marks": array("Q"), "vector_boxes": array("f"), "shape_marks": array("Q"), "hiding_marks": array("Q"), "text_marks": array("Q"),
            "runs_known": False, "marks_known": False, "visual_sha256": None}


def _inspect_page(page: Any, number: int, limits: Limits, memo: dict | None = None, expired: Callable[[], bool] | None = None) -> dict:
    """Text, heading, pictures and drawing of one page.

    Besides the text, the page is read as a list of *marks*: one per painted shape (its path, where it is placed, its
    colour and line), one per placed picture, one per text run (its string, position, font and colour). A mark is a
    digest, so two pages can be compared mark by mark — what is on both, what was added, what is gone — and a page
    gets one digest of everything it draws, in order (``visual_sha256``). Marks inside a form are read in the form's
    own coordinates together with where the form is placed. Nothing is rendered, so what covers what is not known.

    Two kinds of shape are compared like any mark but are not counted as a drawing, because they show nothing of their
    own: a plain box that covers the whole page (a background), and a shape filled and outlined in nothing but the
    paint of that background (or in white on a page that paints no background) — a cover, which can only hide what
    lies under it. Once a picture covers the page, what is painted afterwards lies on a picture and not on paper:
    white paint is then a shape like any other, never a cover.

    Each shape and picture is also given the box it lies in on the page (``vector_boxes``, also for marks inside a
    form), so that what most pages draw in the body of the page can be told from what they carry in their margins.
    A background, a cover and a plain closed box over a fifth of the page or more (a panel, a frame) get no box:
    none of them is part of a figure. Text runs painted in the paint of the background are counted
    (``text_in_background_paint``): the file holds those words, the page does not show them.
    """
    memo = memo if memo is not None else {}
    box = page.mediabox
    width, height, left, bottom = float(box.width), float(box.height), float(box.left), float(box.bottom)
    resources = _deref(page.get("/Resources")) or {}
    counts = {"paint": 0, "path": 0, "text_ops": 0, "operators": 0, "unseen_text": 0}
    placed: list[tuple[str, list[float]]] = []
    chunks: list[tuple[str, float, float, float]] = []
    inline: list[dict] = []
    vector, shapes, hiding, text_marks, boxes = array("Q"), array("Q"), array("Q"), array("Q"), array("f")
    visual = hashlib.sha256()
    marks = {"known": True, "truncated": False, "late": False, "text_in_forms": False}
    state = {"f": "", "fcs": "", "s": "", "scs": "", "l": {}, "gs": "", "font": "", "t": {}}
    saved: list[tuple[str, dict]] = []  # graphics states put aside by q, and by entering a form
    frames: list[dict] = [{"resources": resources, "context": "", "matrix": _IDENTITY}]  # the page, then one frame per form being drawn (with the way from its coordinates to the page's)
    opened: list[dict] = []  # one entry per Do that has not ended
    waiting: list[dict | None] = [None]  # the Do whose first inner operator has not come yet
    path: list[str] = []
    outline: list[tuple[float, float] | None] = []  # the corners of the path being built while it is a plain polygon; None once it is not
    bounds: list[float] = []  # the extent of the path being built, in the coordinates it is written in: least x, least y, greatest x, greatest y
    clip = [False]
    paper = [None]  # the paint of the page's background, once one is painted
    # Where something has been laid over the background so far (a fill in another paint, a picture, a shading): text there is not on
    # bare paper. ``None`` stands for "somewhere, not known where" — text is then never counted as painted on the background.
    grounds: list[list[float] | None] = []
    clipped: list[list[float] | None] = [None]  # where the last clipping path lies (a shading fills it)

    def lookup(category: str, name: Any) -> Any:
        for frame in (frames[-1], frames[0]):  # a form without resources of its own draws with those of the page
            table = _deref((frame["resources"] or {}).get(category))
            if table is not None and name in table:
                return _raw(table, name)
        return None

    def named(category: str, name: Any) -> str:
        found = lookup(category, name)
        return _identity(found, memo) if found is not None else f"?{name}"

    def reach(x: float, y: float) -> None:
        if bounds:
            bounds[:] = (min(bounds[0], x), min(bounds[1], y), max(bounds[2], x), max(bounds[3], y))
        else:
            bounds.extend((x, y, x, y))

    def on_page(cm: Any, extent: Any) -> list[float] | None:
        """Where a box written in the current coordinates lies on the page, as fractions from the top left (``None`` when that is not known)."""
        if not extent or frames[-1]["matrix"] is None:
            return None
        x0, y0, x1, y1 = extent
        return _box_fraction(list(_mul(cm, frames[-1]["matrix"])), [(x0, y0), (x1, y0), (x0, y1), (x1, y1)], left, bottom, width, height)

    def laid(where: list[float] | None) -> None:
        if where is None or len(grounds) >= 4000:
            grounds[:] = [None]
        elif grounds != [None] and _area(where) > 0:
            grounds.append(where)

    def on_bare_background(cm: Any, tm: Any) -> bool:
        """Whether a text run starts where nothing has been laid over the background (as far as the file tells)."""
        if grounds == [None] or frames[-1]["matrix"] is None:
            return False
        spot = on_page(cm, (float(tm[4]), float(tm[5]), float(tm[4]), float(tm[5])))
        if spot is None:
            return False
        x, y = spot[0], spot[1]
        return not any(box[0] - 0.005 <= x <= box[2] + 0.005 and box[1] - 0.005 <= y <= box[3] + 0.005 for box in grounds)

    def emit(kind: str, data: str, where: list[float] | None = None) -> None:
        digest = hashlib.blake2b((kind + frames[-1]["context"] + "|" + data).encode("utf-8", "backslashreplace"), digest_size=8).digest()
        visual.update(digest)
        if kind == "c":  # a clipping path is part of what the page draws, but it is not a shape of its own
            return
        if len(vector) + len(text_marks) >= limits.max_marks_per_page:
            marks["truncated"] = True
            return
        value = int.from_bytes(digest, "big")
        if kind == "t":
            text_marks.append(value)
        else:
            vector.append(value)
            boxes.extend(where if where is not None else _NO_BOX)
            if kind == "v":  # "b" a background, "h" a fill in the background's paint, "p" a placed picture: compared, but not a painted shape
                shapes.append(value)
            elif kind == "h":
                hiding.append(value)

    def restore(marker: str) -> None:
        while saved:
            kind, earlier = saved.pop()
            if kind == marker or marker == "form":
                state.clear()
                state.update(earlier)
            if kind == marker:
                return

    def covers_the_page(cm: Any) -> bool:
        """Whether the path is a plain box (a rectangle, or four straight sides) over nine tenths of the page or more."""
        if len(frames) != 1 or not outline or None in outline or len(path) > 6:
            return False
        fraction = _box_fraction([float(v) for v in cm], outline, left, bottom, width, height)
        return fraction is not None and (fraction[2] - fraction[0]) * (fraction[3] - fraction[1]) >= 0.9

    def track(operator: bytes, operands: Any, cm: Any, tm: Any) -> None:
        name = operator.decode("latin-1")
        if operator == b"q":
            saved.append(("q", {**state, "l": dict(state["l"]), "t": dict(state["t"])}))
        elif operator == b"Q":
            if saved and saved[-1][0] == "q":
                restore("q")
        elif operator in _PATH or operator == b"h":
            counts["path"] += operator in _PATH
            path.append(name + _fmt(operands))
            if operator == b"re" and len(operands) >= 4:
                reach(float(operands[0]), float(operands[1]))
                reach(float(operands[0]) + float(operands[2]), float(operands[1]) + float(operands[3]))
            elif operator != b"h":  # the points of a line and the control points of a curve: the curve lies within them
                for index in range(0, len(operands) - 1, 2):
                    reach(float(operands[index]), float(operands[index + 1]))
            if len(path) <= 6 and None not in outline:
                if operator == b"re" and len(operands) >= 4:
                    x, y, w, h = (float(v) for v in operands[:4])
                    outline.extend(((x, y), (x + w, y), (x, y + h), (x + w, y + h)))
                elif operator in (b"m", b"l") and len(operands) >= 2:
                    outline.append((float(operands[0]), float(operands[1])))
                elif operator != b"h":
                    outline.append(None)  # a curve
        elif operator in (b"W", b"W*"):
            clip[0] = True
            path.append(name)
        elif operator == b"n":
            if clip[0] and path:
                emit("c", _fmt(cm) + "|" + ";".join(path))
            path.clear()
            outline.clear()
            bounds.clear()
            clip[0] = False
        elif operator in _STROKE_OPS or operator in _FILL_OPS or operator in _BOTH_OPS:
            counts["paint"] += 1
            look = state["gs"]
            if operator not in _FILL_OPS:
                look += "|" + state["scs"] + state["s"] + "|" + _fmt(sorted(state["l"].items()))
            if operator not in _STROKE_OPS:
                look += "|" + state["fcs"] + state["f"]
            fill = state["fcs"] + state["f"] if operator not in _STROKE_OPS else None
            line = state["scs"] + state["s"].lower() if operator not in _FILL_OPS else None  # worded as the fill state words the same paint
            kind = "v"
            if operator in _FILL_OPS and covers_the_page(cm):
                kind, paper[0] = "b", fill
            elif all(paint == paper[0] or (paper[0] is None and paint in _WHITE) for paint in (fill, line) if paint is not None):
                kind = "h"  # filled and outlined in nothing but the paint of the background
            where = on_page(cm, bounds) if kind == "v" else None
            if where is not None and len(outline) >= 4 and None not in outline and len(path) <= 6 and _area(where) >= _PANEL and (
                    operator not in _STROKE_OPS or operator == b"s" or path[-1] == "h" or all(item.startswith("re") for item in path)):
                where = None  # a plain closed box over a good part of the page: a panel or a frame, which shows no figure of its own
            emit(kind, f"{name}|{_fmt(cm)}|{look}|" + ";".join(path), where)
            path.clear()
            outline.clear()
            bounds.clear()
            clip[0] = False
        elif operator == b"sh":
            counts["paint"] += 1
            emit("v", f"sh|{_fmt(cm)}|{state['gs']}|" + (named("/Shading", operands[0]) if operands else ""))
        elif operator in _TEXT_SHOW:
            counts["text_ops"] += 1
            if len(frames) > 1:
                marks["text_in_forms"] = True
            ink = state["fcs"] + state["f"]
            if ink == paper[0] or (paper[0] is None and ink in _WHITE):
                counts["unseen_text"] += 1  # painted in the paint of the background: held by the file, not shown by the page
            emit("t", f"{name}|{_fmt(cm)}|{_fmt(tm)}|{state['gs']}|{ink}|{state['font']}|{_fmt(sorted(state['t'].items()))}|{_fmt(operands)}")
        elif operator in _FILL_COLOUR or operator in _STROKE_COLOUR:
            values = list(operands)
            if values and isinstance(values[-1], str) and not isinstance(values[-1], (int, float)):  # a pattern is named, not numbered
                values[-1] = named("/Pattern", values[-1])
            state["f" if operator in _FILL_COLOUR else "s"] = name + _fmt(values)
            if operator in _DEVICE_COLOUR:  # g, rg, k (G, RG, K) also choose the colour space
                state["fcs" if operator in _FILL_COLOUR else "scs"] = ""
        elif operator in (b"cs", b"CS"):
            space = str(operands[0]) if operands else ""
            state["fcs" if operator == b"cs" else "scs"] = space if space.startswith("/Device") or space == "/Pattern" else named("/ColorSpace", operands[0] if operands else "")
        elif operator in _LINE_STATE:
            state["l"][name] = _fmt(operands)
        elif operator in _TEXT_STATE:
            state["t"][name] = _fmt(operands)
        elif operator == b"Tf":
            state["font"] = (named("/Font", operands[0]) if operands else "") + "@" + _fmt(operands[1:])
        elif operator == b"gs":
            state["gs"] = named("/ExtGState", operands[0]) if operands else ""
        elif operator == b"INLINE IMAGE":
            settings = (operands.get("settings") or {}) if isinstance(operands, dict) else {}
            data = operands.get("data") if isinstance(operands, dict) else None
            data = bytes(data) if isinstance(data, (bytes, bytearray)) else b""
            digest = hashlib.sha256(data).hexdigest()
            where = on_page(cm, (0.0, 0.0, 1.0, 1.0))
            if where is not None and _area(where) >= _PAGE_COVER:
                paper[0], where = _GROUND_PICTURE, None
            emit("p", f"inline|{_fmt(cm)}|{digest}", where)
            pixels = [settings.get(key) for key in ("/W", "/Width", "/H", "/Height")]
            inline.append({"name": "inline", "width": next((int(v) for v in pixels[:2] if v is not None), None), "height": next((int(v) for v in pixels[2:] if v is not None), None),
                           "bytes": len(data), "sha256": digest, "digest_basis": "inline_bytes", "template": False,
                           "box": _box_fraction([float(v) for v in cm], [(0, 0), (1, 0), (0, 1), (1, 1)], left, bottom, width, height) if len(frames) == 1 else None})
        elif operator == b"Do":
            ref = lookup("/XObject", operands[0]) if operands else None
            target = _deref(ref) if ref is not None else None
            entry: dict[str, Any] = {"entered": False, "frame": None, "mark": None}
            if isinstance(target, dict) and target.get("/Subtype") == "/Image":
                picture = _image_record(str(operands[0]), ref, None, memo)
                where = on_page(cm, (0.0, 0.0, 1.0, 1.0))
                if where is not None and _area(where) >= _PAGE_COVER:
                    paper[0], where = _GROUND_PICTURE, None  # a picture over the whole page: what is painted after it lies on a picture, not on paper
                stencil = target.get("/ImageMask")  # a stencil is painted in the fill colour of the moment: that colour is part of what the page shows
                emit("p", f"image|{_fmt(cm)}|" + (picture["sha256"] if picture["digest_basis"] != "name_and_size" else f"unknown:{number}:{len(vector)}")
                     + ("|" + state["fcs"] + state["f"] if getattr(stencil, "value", stencil) is True else ""), where)
            elif isinstance(target, dict):
                placement = _fmt(cm) + "|" + _fmt(target.get("/Matrix") or []) + "|" + _fmt(target.get("/BBox") or [])
                inner, whole = None, None  # the way from the form's coordinates to the page's, and where the form as a whole lies
                try:
                    if frames[-1]["matrix"] is not None:
                        inner = _mul(_mul(target.get("/Matrix") or _IDENTITY, cm), frames[-1]["matrix"])
                        x0, y0, x1, y1 = (float(v) for v in target.get("/BBox") or ())
                        whole = _box_fraction(list(inner), [(x0, y0), (x1, y0), (x0, y1), (x1, y1)], left, bottom, width, height)
                except (TypeError, ValueError):  # a matrix or a box that is not six or four numbers: marks inside the form then have no known place
                    whole = None
                entry["frame"] = {"resources": _deref(target.get("/Resources")) or frames[-1]["resources"], "matrix": inner,
                                  "context": hashlib.blake2b((frames[-1]["context"] + "|" + placement).encode("utf-8"), digest_size=6).hexdigest()}
                # the form as a whole, needed only if the parser does not open it
                entry["mark"] = (f"form|{_fmt(cm)}|", ref, whole if whole is not None and _area(whole) < _PAGE_COVER else None)
            if len(frames) == 1 and operands:
                placed.append((str(operands[0]), [float(v) for v in cm]))
            opened.append(entry)
            waiting[0] = entry

    def before(operator: bytes, operands: Any, cm: Any, tm: Any) -> None:
        counts["operators"] += 1
        if marks["late"] or (expired is not None and counts["operators"] % 2048 == 0 and expired()):
            marks["late"] = True  # said again at every operator: the parser swallows what is raised inside a form
            raise _TimeLimit()
        entry = waiting[0]
        if entry is not None:  # the previous Do opened a form: this operator is the first one inside it
            waiting[0] = None
            entry["entered"] = True
            saved.append(("form", {**state, "l": dict(state["l"]), "t": dict(state["t"])}))
            frames.append(entry["frame"] or {"resources": frames[-1]["resources"], "context": frames[-1]["context"] + "?", "matrix": frames[-1]["matrix"]})
        if operator == b"Do" or marks["known"]:
            try:
                track(operator, operands, cm, tm)
            except Exception:  # an operand of a shape this reader did not expect: the page keeps its text, its drawing is not compared
                marks["known"] = False
                if operator == b"Do":
                    opened.append({"entered": False, "frame": None, "mark": None})
                    waiting[0] = opened[-1]

    def after(operator: bytes, operands: Any, cm: Any, tm: Any) -> None:
        if operator != b"Do" or not opened:
            return
        entry = opened.pop()
        if waiting[0] is entry:
            waiting[0] = None
        if entry["entered"]:
            frames.pop()
            restore("form")
        elif entry["mark"] and marks["known"]:
            try:  # a form the parser did not open (no resources, a repeat, too many): one mark for the whole of it
                emit("v", entry["mark"][0] + _identity(entry["mark"][1], memo), entry["mark"][2])
            except Exception:  # its bytes could not be read: the drawing of the page is then not compared
                marks["known"] = False

    def on_text(text: str, cm: Any, tm: Any, font: Any, size: Any) -> None:
        if len(frames) == 1 and text and text.strip():
            scale = math.hypot(float(tm[0]), float(tm[1])) * math.hypot(float(cm[0]), float(cm[1]))
            x = float(tm[4]) * float(cm[0]) + float(tm[5]) * float(cm[2]) + float(cm[4])
            baseline = float(tm[4]) * float(cm[1]) + float(tm[5]) * float(cm[3]) + float(cm[5])
            chunks.append((_clean(text).strip(), round(float(size or 0) * (scale or 1.0), 2), round(baseline, 2), round(x, 2)))

    text = _clean(page.extract_text(visitor_operand_before=before, visitor_operand_after=after, visitor_text=on_text) or "").strip()
    if marks["late"]:
        raise _TimeLimit()
    truncated = len(text) > limits.max_page_text_chars
    if truncated:
        text = text[:limits.max_page_text_chars]

    images: list[dict] = []
    placements: list[dict] = []
    xobjects = _deref(resources.get("/XObject")) or {}
    count = 0
    for name, matrix in placed:
        ref = _raw(xobjects, name) if name in xobjects else None
        obj = _deref(ref) if ref is not None else None
        if obj is None:
            continue
        subtype = obj.get("/Subtype")
        if subtype == "/Image":
            fraction = _box_fraction(matrix, [(0, 0), (1, 0), (0, 1), (1, 1)], left, bottom, width, height)
            count += 1
            if len(images) < limits.max_images_per_page:
                images.append(_image_record(name, ref, fraction, memo))
            placements.append({"kind": "image", "box": fraction})
        elif subtype == "/Form":
            bbox = [float(v) for v in (obj.get("/BBox") or [0, 0, 1, 1])]
            inner = [float(v) for v in (obj.get("/Matrix") or [1, 0, 0, 1, 0, 0])]
            corners = [(inner[0] * x + inner[2] * y + inner[4], inner[1] * x + inner[3] * y + inner[5]) for x, y in ((bbox[0], bbox[1]), (bbox[2], bbox[1]), (bbox[0], bbox[3]), (bbox[2], bbox[3]))]
            fraction = _box_fraction(matrix, corners, left, bottom, width, height)
            nested: list[dict] = []
            _form_images(obj, fraction, 1, nested, limits.max_images_per_page, memo)
            count += len(nested)
            images.extend(nested[:max(0, limits.max_images_per_page - len(images))])
            placements.append({"kind": "form", "box": fraction})
    count += len(inline)  # pictures written into the page description itself
    images.extend(inline[:max(0, limits.max_images_per_page - len(images))])
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    return {"page": number, "status": "ok", "reason": None, "text": text, "text_truncated": truncated, "word_count": len(_WORD.findall(text)), "title": _title_from(chunks, lines),
            "images": images, "image_count": count, "images_truncated": count > len(images), "paint_operations": counts["paint"], "path_operations": counts["path"],
            "text_operations": counts["text_ops"], "text_in_background_paint": counts["unseen_text"] if marks["known"] else None,
            "size": {"width_pt": round(width, 2), "height_pt": round(height, 2), "rotation": int(page.get("/Rotate") or 0) % 360}, "origin": (left, bottom),
            "placements": placements, "chunks": chunks, "vector_marks": vector, "vector_boxes": boxes, "shape_marks": shapes, "hiding_marks": hiding, "text_marks": text_marks,
            "runs_known": not marks["text_in_forms"] and len(chunks) <= limits.max_marks_per_page,  # text inside a form has no place on the page here
            "marks_known": marks["known"] and not marks["truncated"], "visual_sha256": visual.hexdigest()[:32] if marks["known"] else None}


_TRANSIENT_REASONS = frozenset({"read_error", "time_limit"})  # failures of the moment, not properties of the bytes


def _inspect_pdf(path: Path, limits: Limits, clock: Callable[[], float], started: float | None = None) -> dict:
    """Every page of a PDF, within the bounds. A page that cannot be read stays in the list as ``unknown`` with the
    reason: ``unreadable`` (the parser cannot read these bytes), ``read_error`` (the file could not be read just now),
    ``time_limit`` or ``page_limit``. ``transient`` says that another attempt may give more."""
    result: dict[str, Any] = {"status": "complete", "reason": None, "transient": False, "page_count": None, "pages": [], "labels_defined": False,
                              "tool": f"pypdf {metadata.version('pypdf')}"}
    try:
        size = path.stat().st_size
    except OSError:
        return {**result, "status": "unavailable", "reason": "missing"}
    if size == 0:
        return {**result, "status": "unavailable", "reason": "empty"}
    if size > limits.max_bytes:
        return {**result, "status": "not_inspected", "reason": "byte_limit"}
    started = clock() if started is None else started
    try:
        reader = PdfReader(str(path))
        if reader.is_encrypted and not reader.decrypt(""):
            return {**result, "status": "unreadable", "reason": "encrypted"}
        count = len(reader.pages)
    except OSError:  # locked, or a hiccup of a streamed folder: says nothing about the bytes
        return {**result, "status": "unavailable", "reason": "read_error", "transient": True}
    except Exception as exc:  # pypdf raises many types on damaged files; none of them may escape as a crash
        return {**result, "status": "unreadable", "reason": type(exc).__name__}
    labels: list[str] | None = None
    try:
        if "/PageLabels" in _deref(reader.trailer["/Root"]):
            labels = [str(label) for label in reader.page_labels]
    except Exception:
        labels = None
    result.update(page_count=count, labels_defined=labels is not None and len(labels) == count)
    memo: dict = {}
    late = False

    def expired() -> bool:
        return clock() - started > limits.max_seconds

    for index in range(count):
        number = index + PAGE_INDEX_BASE
        if number > limits.max_pages:
            record = _unknown_page(number, "page_limit")
        elif late or expired():
            late = True
            record = _unknown_page(number, "time_limit")
        else:
            try:
                record = _inspect_page(reader.pages[index], number, limits, memo, expired)
            except _TimeLimit:
                late = True
                record = _unknown_page(number, "time_limit")
            except OSError:
                record = _unknown_page(number, "read_error")
            except Exception:
                record = _unknown_page(number, "unreadable")
        record["pdf_label"] = labels[index] if result["labels_defined"] else None
        result["pages"].append(record)
    if any(page["status"] != "ok" for page in result["pages"]):
        result["status"] = "partial"
        result["transient"] = any(page["reason"] in _TRANSIENT_REASONS for page in result["pages"])
    return result


# ---- PPTX inspection (python-pptx: a parser, never a renderer) -----------------------------------------------------
def _shape_texts(shapes: Any, out: list[str], depth: int = 0) -> None:
    for shape in shapes:
        if getattr(shape, "has_text_frame", False) and shape.text_frame is not None:
            text = shape.text_frame.text
            if text and text.strip():
                out.append(text)
        if getattr(shape, "has_table", False):
            for row in shape.table.rows:
                for cell in row.cells:
                    if cell.text and cell.text.strip():
                        out.append(cell.text)
        inner = getattr(shape, "shapes", None)
        if inner is not None and depth < 4:
            _shape_texts(inner, out, depth + 1)


def inspect_pptx(path: Path, limits: Limits | None = None, *, clock: Callable[[], float] | None = None, started: float | None = None) -> dict[str, Any]:
    """Slide id, order, title, text, hidden flag and layout name of a ``.pptx`` file. Identity and text only.

    Never raises for a damaged or oversized file: the state is reported (``unreadable`` / ``not_inspected``). The size
    is checked before a byte is read. The time bound (``limits.max_seconds``, counted from ``started`` when a caller
    has already spent time) is checked before the file is opened and between slides: a presentation that runs out of
    time keeps the slides read so far and says ``truncated`` with the reason ``time_limit``. Opening the file is one
    call into the parser and cannot be interrupted.
    """
    limits = limits or Limits()
    clock = clock or time.monotonic
    result: dict[str, Any] = {"status": "ok", "reason": None, "slide_count": None, "slides": [], "hidden_slides": [], "truncated": False,
                              "tool": f"python-pptx {metadata.version('python-pptx')}"}
    try:
        size = path.stat().st_size
    except OSError:
        return {**result, "status": "unavailable", "reason": "missing"}
    if path.suffix.lower() != ".pptx":
        return {**result, "status": "not_inspected", "reason": "not_pptx"}
    if size > limits.max_pptx_bytes:
        return {**result, "status": "not_inspected", "reason": "byte_limit"}
    if started is None:
        started = clock()
    elif clock() - started > limits.max_seconds:
        return {**result, "status": "not_inspected", "reason": "time_limit"}
    try:
        with zipfile.ZipFile(path) as archive:
            if sum(info.file_size for info in archive.infolist()) > limits.max_pptx_expanded_bytes:
                return {**result, "status": "not_inspected", "reason": "expanded_size_limit"}
        from pptx import Presentation  # imported here: only a unit that has a presentation pays for the parser

        deck = Presentation(str(path))
        slides = []
        for ordinal, slide in enumerate(deck.slides, start=1):
            if ordinal > limits.max_pptx_slides:
                result.update(truncated=True, reason="slide_limit")
                break
            if clock() - started > limits.max_seconds:
                result.update(truncated=True, reason="time_limit")
                break
            texts: list[str] = []
            _shape_texts(slide.shapes, texts)
            title_shape = slide.shapes.title
            title = " ".join(title_shape.text_frame.text.split()) if title_shape is not None and title_shape.has_text_frame else ""
            text = _clean("\n".join(texts).replace("\x0b", "\n")).strip()[:limits.max_page_text_chars]
            slides.append({"slide_id": int(slide.slide_id), "ordinal": ordinal, "title": title or None, "text": text, "word_count": len(_WORD.findall(text)),
                           "hidden": slide.element.get("show") in {"0", "false"}, "layout": slide.slide_layout.name})
    except OSError:  # locked, or a hiccup of a streamed folder: says nothing about the bytes
        return {**result, "status": "unavailable", "reason": "read_error"}
    except Exception as exc:  # not a presentation, or one the parser cannot read
        if type(exc).__name__ == "PackageNotFoundError":  # the archive was opened a moment ago: the parser not finding it now is a state of the moment
            return {**result, "status": "unavailable", "reason": "read_error"}
        return {**result, "status": "unreadable", "reason": type(exc).__name__}
    result.update(slide_count=len(deck.slides), slides=slides, hidden_slides=[s["ordinal"] for s in slides if s["hidden"]])
    return result


# ---- analysis ----------------------------------------------------------------------------------------------------
@dataclass
class _Feat:
    page: int
    ok: bool
    lines: list[str] = field(default_factory=list)
    line_tokens: list[tuple[str, ...]] = field(default_factory=list)
    tokens: Counter = field(default_factory=Counter)
    words: Counter = field(default_factory=Counter)
    body: Counter = field(default_factory=Counter)
    title: str | None = None
    title_norm: str = ""
    images: Counter = field(default_factory=Counter)  # pictures that are not part of the template, by what identifies them
    pictures_unread: bool = False  # a picture whose bytes could not be read: it is never taken to equal another
    marks: Counter = field(default_factory=Counter)  # painted shapes, placed pictures and backgrounds beyond the template: what the page shows
    covers: Counter = field(default_factory=Counter)  # fills in the paint of the background, beyond the template: they show nothing and can only hide
    text_marks: Counter = field(default_factory=Counter)  # text runs as they are painted: string, place, font and colour
    runs: Counter = field(default_factory=Counter)  # text runs by where they stand: string, baseline, left edge, size
    runs_known: bool = False  # every text run of the page has a place (none is drawn inside a form)
    shapes: int = 0  # painted shapes beyond the template (a page-covering fill is a background, not a shape)
    marks_known: bool = False  # the drawing of the page was read mark by mark and can be compared
    visual: str | None = None  # one digest of everything the page draws, in order
    painted: int = 0
    beyond_first_line: int = 0  # words after the first line: how much text the page holds, whatever was taken for its heading
    portrait: bool = False
    printed: int | None = None
    label: str | None = None
    pptx_id: int | None = None

    @property
    def total(self) -> int:
        return sum(self.tokens.values())

    @property
    def word_total(self) -> int:
        return sum(self.words.values())

    @property
    def body_total(self) -> int:
        return sum(self.body.values())


def _is_portrait(size: dict | None) -> bool:
    if not size:
        return False
    width, height = (size["height_pt"], size["width_pt"]) if size.get("rotation") in (90, 270) else (size["width_pt"], size["height_pt"])
    return height > width * 1.1


def _number_candidates(page: dict) -> list[tuple[int, float, float]]:
    """Short isolated numbers in the top or bottom margin of a page: ``(value, from the top, from the left)`` as shares of the page."""
    size, origin = page.get("size"), page.get("origin")
    if not size or not origin or size["height_pt"] <= 0 or size["width_pt"] <= 0:
        return []
    found = []
    for text, _, baseline, x in page.get("chunks") or []:
        match = _NUMBER_RUN.match(text)
        if not match:
            continue
        from_top, from_left = 1 - (baseline - origin[1]) / size["height_pt"], (x - origin[0]) / size["width_pt"]
        if (0 <= from_top <= _MARGIN_TOP or _MARGIN_BOTTOM <= from_top <= 1) and 0 <= from_left <= 1:
            found.append((int(match.group(1)), round(from_top, 4), round(from_left, 4)))
    return found[:8]


def _detect_printed_numbers(pages: list[dict]) -> tuple[dict[int, int], dict[str, Any]]:
    """Printed slide numbers, only when one place in the page margin carries a consistent sequence across the deck.

    Two kinds of evidence are needed. Position: a short text run that is nothing but a number, standing at the same
    place in the top or bottom margin of most pages — a number that ends a sentence, a caption or a list is body text
    and is never read as a slide number. Sequence: at least three different numbers that never run backwards and
    mostly advance no faster than the pages do (a number may repeat on the builds of one slide, and a few slides may
    be left out). Where the number repeats, it must mostly do so under one heading, as on the builds of a slide: a
    counter that stays the same over different slides (a section, a part) is not a slide number. A lone number, a
    constant, or numbers that jump about establish nothing.
    """
    nothing = {"method": "consistent_number_sequence", "result": "no place in the page margin carries a consistent sequence of isolated numbers"}
    candidates = {page["page"]: _number_candidates(page) for page in pages if page["status"] == "ok"}
    candidates = {number: found for number, found in candidates.items() if found}
    text_pages = sum(1 for page in pages if page["status"] == "ok" and page["lines"])
    headings = {page["page"]: _norm(page.get("title") or "") for page in pages}
    if len(candidates) < 3:
        return {}, nothing
    best: tuple[dict[int, int], dict[str, Any]] | None = None
    anchors: set[tuple[float, float]] = set()
    for found in candidates.values():
        for _, top, left in found:
            anchor = (round(top, 2), round(left, 2))
            if anchor in anchors or len(anchors) >= 200:
                continue
            anchors.add(anchor)
            entries: list[tuple[int, int, float, float]] = []  # page, value, from top, from left — the number nearest to this place on each page
            for number, others in sorted(candidates.items()):
                near = [(abs(t - top) + abs(x - left), value, t, x) for value, t, x in others if abs(t - top) <= 0.03 and abs(x - left) <= 0.08]
                if near:
                    _, value, t, x = min(near)
                    entries.append((number, value, t, x))
            if len(entries) < 3 or (best is not None and len(entries) <= len(best[0])):
                continue
            # Longest chain that never decreases and never advances much faster than the pages do.
            length, back = [1] * len(entries), [-1] * len(entries)
            for j, (page_j, value_j, _, _) in enumerate(entries):
                for i in range(max(0, j - 200), j):
                    page_i, value_i = entries[i][0], entries[i][1]
                    if 0 <= value_j - value_i <= (page_j - page_i) + 25 and length[i] + 1 > length[j]:
                        length[j], back[j] = length[i] + 1, i
            end = max(range(len(entries)), key=lambda k: (length[k], -k))
            chain = []
            while end != -1:
                chain.append(entries[end])
                end = back[end]
            chain.reverse()
            distinct = len({value for _, value, _, _ in chain})
            steps = [(a[1], b[1], b[0] - a[0]) for a, b in zip(chain, chain[1:])]
            unit = sum(1 for first, second, apart in steps if second - first <= apart)  # the number advances no faster than the pages
            if len(chain) < 3 or distinct < 3 or len(chain) < 0.6 * text_pages or len(chain) < 0.85 * len(entries) or unit < 0.6 * len(steps):
                continue
            # A slide number repeats on the builds of one slide: pages under one heading. A number that mostly repeats over pages with
            # different headings counts something else (a section, a part) and is not worded as a slide number.
            repeated = [(a[0], b[0]) for a, b in zip(chain, chain[1:]) if a[1] == b[1]]
            one_heading = sum(1 for first, second in repeated if headings.get(first) and headings[first] == headings.get(second))
            if repeated and one_heading < 0.5 * len(repeated):
                continue
            if best is not None and len(chain) <= len(best[0]):
                continue
            tops, lefts = sorted(t for _, _, t, _ in chain), sorted(x for _, _, _, x in chain)
            evidence = {"method": "consistent_number_sequence", "position": "page_margin", "margin": "top" if tops[len(tops) // 2] <= 0.5 else "bottom",
                        "from_top": tops[len(tops) // 2], "from_left": lefts[len(lefts) // 2], "pages_with_number": len(chain), "pages_with_text": text_pages,
                        "candidates": len(entries), "distinct_values": distinct, "first": [chain[0][0], chain[0][1]], "last": [chain[-1][0], chain[-1][1]],
                        "repeats": len(repeated), "repeats_under_one_heading": one_heading, "steps_of_at_most_one": unit,
                        "skips": [[first, second] for first, second, apart in steps if second - first > apart][:50],
                        "rejected_pages": sorted({entry[0] for entry in entries} - {entry[0] for entry in chain})[:50]}
            best = ({entry[0]: entry[1] for entry in chain}, evidence)
    return best if best is not None else ({}, nothing)


def _strip_number(lines: list[str], value: int, margin: str | None) -> list[str]:
    """The page's lines without its printed number (so that the builds of one slide compare equal in a deck that numbers every page)."""
    lines = list(lines)
    order = range(len(lines)) if margin == "top" else range(len(lines) - 1, -1, -1)
    for index in order:  # the number as a line of its own
        match = _NUMBER_RUN.match(lines[index])
        if match and int(match.group(1)) == value:
            del lines[index]
            return lines
    for index in dict.fromkeys((len(lines) - 1, 0)):  # the number at the end of a line it shares with a footer
        if 0 <= index < len(lines):
            match = _TRAILING_NUMBER.search(lines[index])
            if match and int(match.group(1)) == value:
                lines[index] = lines[index][:match.start(1)].rstrip(" -–—|·•/")
                return [line for line in lines if line.strip()]
    return lines


def _text_runs(page: dict, printed: int | None) -> Counter:
    """The text runs of a page by where they stand — string, baseline, left edge and size, to the point — whatever colour, font object or
    order they are painted in. The printed number is left out: in a deck that numbers every page it differs between the builds of one slide.
    The offset keeps half and whole points away from the edge of a cell."""
    runs: Counter = Counter()
    for string, size, baseline, x in page["chunks"]:
        if printed is not None:
            match = _NUMBER_RUN.match(string)
            if match and int(match.group(1)) == printed:
                continue
        runs[(_norm(string), math.floor(baseline + 0.37), math.floor(x + 0.37), round(size))] += 1
    return runs


def _features(pages: list[dict], printed: dict[int, int], margin: str | None) -> tuple[list[_Feat], dict[str, Any]]:
    ok = [page for page in pages if page["status"] == "ok"]
    stripped = {page["page"]: (_strip_number(page["lines"], printed[page["page"]], margin) if page["page"] in printed else list(page["lines"])) for page in ok}
    titles = {page["page"]: page["title"] for page in ok}
    # Lines repeated on most pages (a footer, a running header) carry no page content. A heading is never one of them.
    boilerplate: set[str] = set()
    with_text = [page for page in ok if stripped[page["page"]]]
    if len(with_text) >= 6:
        counts: Counter = Counter()
        for page in with_text:
            counts.update({re.sub(r"\d+", "#", _norm(line)) for line in stripped[page["page"]]})
        heading_keys = {re.sub(r"\d+", "#", _norm(title)) for title in titles.values() if title}
        boilerplate = {line for line, n in counts.items() if n >= math.ceil(0.7 * len(with_text)) and line not in heading_keys}
    # Pictures and shapes drawn on nearly every page belong to the template (a logo, a rule under the heading, a background), not to
    # a slide — when they stand in the margins of the page, cover it, or are a plain panel. What nearly every page draws in the BODY
    # of the page may as well be one figure that the deck walks through: it stays content, so that a page showing it is never read
    # as blank or as a bare heading. A picture is a figure from a fiftieth of the page; shapes are, when together they span that.
    template_images: set[str] = set()
    template_marks: set[int] = set()
    body_images: set[str] = set()
    body_marks: set[int] = set()
    if len(ok) >= 5:
        seen: Counter = Counter()
        drawn: Counter = Counter()
        for page in ok:
            seen.update({image["sha256"] for image in page["images"] if image["digest_basis"] != "name_and_size"})
            drawn.update(set(page["vector_marks"]))
        needed = math.ceil(0.8 * len(ok))
        frequent_images = {digest for digest, n in seen.items() if n >= needed}
        frequent_marks = {mark for mark, n in drawn.items() if n >= needed}
        picture_at: dict[str, Any] = {}
        mark_at: dict[int, Any] = {}
        for page in ok:  # where each of them lies: the same mark lies at the same place on every page that draws it
            if len(picture_at) == len(frequent_images) and len(mark_at) == len(frequent_marks):
                break
            for image in page["images"]:
                if image["sha256"] in frequent_images and image["digest_basis"] != "name_and_size":
                    picture_at.setdefault(image["sha256"], image["box"])
            placed = page["vector_boxes"]
            for index, mark in enumerate(page["vector_marks"]):
                if mark in frequent_marks and mark not in mark_at:
                    mark_at[mark] = tuple(placed[4 * index:4 * index + 4])
        body_images = {digest for digest in frequent_images if _in_body(picture_at.get(digest)) and _area(picture_at[digest]) >= _FIGURE}
        body_marks = {mark for mark in frequent_marks if _in_body(mark_at.get(mark))}
        if body_marks:
            spans = [mark_at[mark] for mark in body_marks]
            if _area((min(s[0] for s in spans), min(s[1] for s in spans), max(s[2] for s in spans), max(s[3] for s in spans))) < _FIGURE:
                body_marks = set()  # a small mark in the middle of every page (an emblem, a divider line) is still the template
        template_images = frequent_images - body_images
        template_marks = frequent_marks - body_marks
    feats: list[_Feat] = []
    for page in pages:
        if page["status"] != "ok":
            feats.append(_Feat(page["page"], False, label=page.get("pdf_label")))
            continue
        lines = [line for line in stripped[page["page"]] if re.sub(r"\d+", "#", _norm(line)) not in boilerplate]
        text = "\n".join(lines)
        title = page["title"]
        if title and page["page"] in printed and _NUMBER_RUN.match(title.strip()):
            title = None  # the printed number itself was the largest text: not a heading
        if title is None or re.sub(r"\d+", "#", _norm(title)) in boilerplate:
            title = next((line for line in lines if sum(ch.isalpha() for ch in line) >= 2), None)
        words = Counter(_words(text))
        body = words - Counter(_words(title or ""))
        for image in page["images"]:
            image["template"] = image["sha256"] in template_images
        pictures = [(index, image) for index, image in enumerate(page["images"]) if not image["template"]]
        line_tokens = [tokens for tokens in (tuple(_TOKEN.findall(_norm(line))) for line in lines) if tokens]
        covers = Counter(mark for mark in page["hiding_marks"] if mark not in template_marks)
        feats.append(_Feat(page["page"], True, lines=lines, line_tokens=line_tokens, tokens=Counter(token for tokens in line_tokens for token in tokens), words=words, body=body,
                           title=title, title_norm=_norm(title or ""), images=Counter(_picture_key(image, page["page"], index) for index, image in pictures),
                           pictures_unread=any(image["digest_basis"] == "name_and_size" for _, image in pictures),
                           marks=Counter(mark for mark in page["vector_marks"] if mark not in template_marks) - covers, covers=covers, text_marks=Counter(page["text_marks"]),
                           runs=_text_runs(page, printed.get(page["page"])), runs_known=bool(page["runs_known"]),
                           shapes=sum(1 for mark in page["shape_marks"] if mark not in template_marks), marks_known=bool(page["marks_known"]), visual=page["visual_sha256"],
                           painted=page["paint_operations"] or 0, beyond_first_line=sum(words.values()) - len(_words(lines[0])) if lines else 0,
                           portrait=_is_portrait(page["size"]), printed=printed.get(page["page"]), label=page.get("pdf_label")))
    return feats, {"boilerplate_lines": len(boilerplate), "template_images": len(template_images), "template_marks": len(template_marks),
                   "repeated_in_body": {"images": len(body_images), "marks": len(body_marks)}}


def _align_pptx(feats: list[_Feat], slides: list[dict]) -> tuple[dict[int, dict], dict[str, Any]]:
    """Monotone sequence alignment of PDF pages to presentation slides on title and text.

    Several consecutive pages may match one slide (builds); pages and slides may stay unmatched. A page with no text
    has nothing to align on and is left unmatched — its position between matched neighbours is not used. Fewer than
    three matched pages in order (or matches that all fall on one slide of a longer presentation) are a coincidence
    of words, not a sequence: the alignment is then ``not_supported`` and no page is given a slide.
    """
    rows, cols = len(feats), len(slides)
    if rows * cols > _MAX_ALIGNMENT_CELLS:  # bounded: a book against a long deck is not compared cell by cell
        return {}, {"status": "not_attempted", "method": "text_title_sequence_alignment", "reason": "too_large", "matched_pages": 0, "candidate_matches": 0,
                    "unmatched_pages": [feat.page for feat in feats], "unmatched_slides": [slide["ordinal"] for slide in slides], "page_equals_ordinal_everywhere": False}
    slide_words = [Counter(_words(slide["text"])) for slide in slides]
    slide_titles = [_norm(slide["title"] or "") for slide in slides]
    score: list[list[float | None]] = [[None] * cols for _ in range(rows)]
    detail: dict[tuple[int, int], dict] = {}
    for i, feat in enumerate(feats):
        total = feat.word_total
        if not feat.ok or total < 2:
            continue
        for j in range(cols):
            common = sum((feat.words & slide_words[j]).values())
            if not common:
                continue
            contained = common / total
            union = len(set(feat.words) | set(slide_words[j]))
            jaccard = len(set(feat.words) & set(slide_words[j])) / union if union else 0.0
            title_equal = bool(feat.title_norm) and feat.title_norm == slide_titles[j]
            if (contained >= 0.8 and (total >= 3 or title_equal)) or (title_equal and contained >= 0.6):
                score[i][j] = contained + 0.5 * jaccard + (0.5 if title_equal else 0.0)
                detail[(i, j)] = {"title_equal": title_equal, "page_text_in_slide": round(contained, 3), "word_overlap": round(jaccard, 3)}
    minus = float("-inf")
    best = [[0.0] * (cols + 1) for _ in range(rows + 1)]  # best[i][j]: first i pages against first j slides
    matched = [[minus] * (cols + 1) for _ in range(rows + 1)]  # ... with page i matched to slide j
    for i in range(1, rows + 1):
        for j in range(1, cols + 1):
            value = score[i - 1][j - 1]
            if value is not None:
                matched[i][j] = value + max(best[i - 1][j - 1], matched[i - 1][j])
            best[i][j] = max(best[i - 1][j], best[i][j - 1], matched[i][j])
    pairs: dict[int, int] = {}
    i, j, on_match = rows, cols, False
    while i > 0 and j > 0:
        if on_match or (matched[i][j] != minus and best[i][j] == matched[i][j]):
            pairs[i - 1] = j - 1
            value = score[i - 1][j - 1] or 0.0
            stay = matched[i - 1][j] != minus and abs(matched[i][j] - (value + matched[i - 1][j])) < 1e-9 and matched[i - 1][j] >= best[i - 1][j - 1]
            i, on_match = i - 1, stay
            if not stay:
                j -= 1
        elif best[i][j] == best[i - 1][j]:
            i -= 1
        else:
            j -= 1
    text_pages = sum(1 for feat in feats if feat.ok and feat.word_total >= 2)
    supported = len(pairs) >= _MIN_ALIGNED_PAGES and len(pairs) >= 0.6 * text_pages and (len(set(pairs.values())) >= 2 or cols < 2)
    result: dict[int, dict] = {}
    if supported:
        for i, j in sorted(pairs.items()):
            result[feats[i].page] = {"slide_id": slides[j]["slide_id"], "ordinal": slides[j]["ordinal"], "evidence": {"method": "text_title_sequence_alignment", **detail[(i, j)]}}
    matched_ordinals = {entry["ordinal"] for entry in result.values()}
    stats = {"status": "supported" if supported else "not_supported", "method": "text_title_sequence_alignment", "matched_pages": len(result), "pages_with_text": text_pages,
             "candidate_matches": len(pairs), "minimum_matched_pages": _MIN_ALIGNED_PAGES,
             "unmatched_pages": [feat.page for feat in feats if feat.page not in result] if supported else [feat.page for feat in feats],
             "unmatched_slides": [slide["ordinal"] for slide in slides if slide["ordinal"] not in matched_ordinals],
             "page_equals_ordinal_everywhere": supported and len(result) == len(feats) and all(entry["ordinal"] == page for page, entry in result.items())}
    return result, stats


def _lines_in_order(first: list[tuple[str, ...]], second: list[tuple[str, ...]]) -> tuple[int, int]:
    """How many lines of the earlier page recur on the later page in the same order — as the same line, or as a line
    that only grew at its end — and how many of those grew. All of them: the later page contains the earlier text."""
    kept = extended = position = 0
    for line in first:
        for index in range(position, len(second)):
            other = second[index]
            if other == line or (len(other) > len(line) and other[:len(line)] == line):
                kept += 1
                extended += other != line
                position = index + 1
                break
    return kept, extended


def _pair(a: _Feat, b: _Feat, labels_signal: bool, numbers_count_slides: bool = False) -> dict | None:
    """What two adjacent pages are to each other: ``join`` (one logical slide), ``uncertain`` (kept separate, with the
    reasons) or ``distinct`` (a shared heading or picture that is not the same slide). ``None``: nothing in common.

    A join is one of three kinds. ``incremental_build``: the later page shows all that the earlier page shows — its
    lines in the same order (or its text runs at the same places), its pictures and its shapes — and more: something
    is added, or a cover is lifted. ``duplicate``: the two pages draw exactly the same, mark for mark.
    ``animation_steps``: one authored slide in two states — the text is the same and most of what is drawn is the
    same, but something changed, moved or was covered; or the drawings could not be compared; or the deck itself
    gives both pages one number.

    Equal words do not carry a join on their own: the same words in another arrangement are never a build, and the
    same heading and caption over another picture, or over a drawing that is mostly different, stay two pages.
    Pages on which the deck prints different slide numbers are not joined when its numbers count slides.
    """
    if not (a.ok and b.ok):
        return None
    ta, tb = a.total, b.total
    common = sum((a.tokens & b.tokens).values())
    removed, added = ta - common, tb - common
    same_title = bool(a.title_norm) and a.title_norm == b.title_norm
    kept, extended = _lines_in_order(a.line_tokens, b.line_tokens)
    compared = a.marks_known and b.marks_known
    placed = a.runs_known and b.runs_known and bool(a.runs)
    # Every text run of the earlier page is on the later page, where it was: painted the same way, or the same string at the same place.
    in_place = (compared and bool(a.text_marks) and not (a.text_marks - b.text_marks)) or (placed and not (a.runs - b.runs))
    same_runs = (compared and bool(a.text_marks) and a.text_marks == b.text_marks) or (placed and a.runs == b.runs)
    flat_a, flat_b = [token for line in a.line_tokens for token in line], [token for line in b.line_tokens for token in line]
    identical = ta > 0 and (flat_a == flat_b or (same_runs and removed == 0 and added == 0))
    superset = ta > 0 and not identical and removed == 0 and added > 0 and (kept == len(a.line_tokens) or in_place)
    same_words = ta > 0 and removed == 0 and not identical and not superset  # every word recurs, but not as the same lines in the same order
    shrinks = tb > 0 and removed > 0 and added == 0
    body_overlap = sum((a.body & b.body).values()) / a.body_total if a.body_total else None
    unread = a.pictures_unread or b.pictures_unread
    images_same = not unread and a.images == b.images
    images_kept = not unread and not (a.images - b.images)
    # What is drawn: shapes and placed pictures on one side, covers (fills in the background's paint, which show nothing) on the other.
    shapes = [sum(a.marks.values()), sum(b.marks.values())]
    shared = sum((a.marks & b.marks).values())
    covers = [sum(a.covers.values()), sum(b.covers.values())]
    cover_removed, cover_added = compared and bool(a.covers - b.covers), compared and bool(b.covers - a.covers)
    drawing_same = compared and a.marks == b.marks and a.covers == b.covers
    # Kept: every shape and picture of the earlier page is on the later page. A cover is not content. One that is lifted shows more. One that is
    # added may hide something of the earlier page: the later page is then not known to show all of it, so the pair is never called a build.
    drawing_kept = compared and not (a.marks - b.marks)
    content_added = drawing_kept and shapes[1] > shapes[0]
    drawing_grows = not cover_added and (content_added or (drawing_kept and cover_removed))
    # Equal marks are sound evidence that two pages draw the same. Different marks are weak evidence that they do not: the same shapes
    # are written with other operators when a page is exported under a transparency group or laid out again.
    mostly_same = compared and (max(shapes) == 0 or shared >= 0.5 * max(shapes))
    pictures_vouch = bool(a.images) and images_kept  # the pictures of the earlier page, compared by their bytes, are all on the later page
    page_identical = not unread and a.visual is not None and a.visual == b.visual
    number_repeat = a.printed is not None and a.printed == b.printed
    numbers_differ = numbers_count_slides and a.printed is not None and b.printed is not None and a.printed != b.printed
    label_repeat = labels_signal and a.label is not None and a.label == b.label
    pptx_same = a.pptx_id is not None and a.pptx_id == b.pptx_id
    authored_same = number_repeat or label_repeat or pptx_same

    signals = ["adjacent"]
    for flag, name in ((same_title, "same_title"), (identical, "identical_text"), (superset, "text_superset"), ((identical or superset) and in_place, "text_in_place"),
                       (same_words, "same_words_other_lines"), (page_identical and (ta > 0 or bool(a.images) or shapes[0] > 0), "page_identical"),
                       (bool(a.images) and images_same, "images_identical"), (bool(a.images) and images_kept and not images_same, "images_contained"),
                       (number_repeat, "printed_number_repeat"), (numbers_differ, "printed_numbers_differ"), (label_repeat, "pdf_label_repeat"), (pptx_same, "pptx_same_slide"),
                       (drawing_same and (shapes[0] > 0 or covers[0] > 0), "drawing_identical"), (drawing_kept and not drawing_same, "drawing_contained"),
                       (compared and not drawing_kept, "drawing_changed"), (cover_removed, "cover_removed"), (cover_added, "cover_added"),
                       (not compared, "drawing_not_compared"), (unread, "picture_not_read")):
        if flag:
            signals.append(name)
    notes = []
    if a.printed is not None and b.printed is not None and a.printed != b.printed:
        notes.append(f"the deck prints different numbers on the two pages ({a.printed}, {b.printed})")
    if a.pptx_id is not None and b.pptx_id is not None and a.pptx_id != b.pptx_id:
        notes.append("the pages align with different presentation slides")
    measures = {"tokens": [ta, tb], "carried_over": round(common / ta, 3) if ta else None, "added": added, "removed": removed,
                "lines": [len(a.line_tokens), len(b.line_tokens)], "lines_kept": kept, "lines_extended": extended,
                "body_overlap": round(body_overlap, 3) if body_overlap is not None else None, "pictures": [sum(a.images.values()), sum(b.images.values())],
                "drawing": {"shapes": shapes, "common": shared} if compared else None, "covers": covers if compared else None}

    def out(decision: str, relation: str | None = None, kind: str | None = None, confidence: str | None = None, reasons: list[str] | None = None) -> dict:
        if decision == "join" and numbers_differ:  # the deck numbers slides and gives these two pages different numbers: its word stands against the reading
            return {"pages": [a.page, b.page], "decision": "uncertain", "relation": None, "kind": None, "confidence": None, "signals": signals,
                    "reasons": [f"the deck prints different slide numbers on the two pages ({a.printed}, {b.printed}); without them the pages would read as {relation}"],
                    "notes": notes, "measures": measures}
        return {"pages": [a.page, b.page], "decision": decision, "relation": relation, "kind": kind, "confidence": confidence, "signals": signals,
                "reasons": reasons or [], "notes": notes, "measures": measures}

    def step(relation: str) -> dict:  # one authored slide in two states: never called a build or a copy, never more than medium
        if relation == "identical_text_visual_change":
            # Named for what the file shows. The same picture on screen can be written in other operations, so a difference on screen is not claimed.
            notes.append("the two pages are drawn differently in the file (other operations, paints or covers); that they look different on screen was not established")
        return out("join", relation, "animation_steps", "medium")

    differs = f"most of the drawing differs ({shared} of {max(shapes)} shapes are on both pages)"
    # A caption's worth of text (a heading and a few labels) can stand over two different figures; a slide's worth of the same text does not.
    ample = a.beyond_first_line >= _AMPLE_TEXT
    # More than a heading: text beyond it, or — on a page whose largest text is its body, so that no heading stands apart — a slide's worth after the first line.
    more_than_a_heading = a.body_total >= 3 or ample
    # What lets equal or growing text carry a join although most of the drawing reads as different.
    carried = mostly_same or ample or pictures_vouch or authored_same
    if ta == 0 and tb == 0:
        content = bool(a.images) or shapes[0] > 0
        if page_identical and content:
            return out("join", "identical", "duplicate", "high")
        if authored_same:
            return step("same_authored_slide")
        if content and images_kept and (drawing_kept or not compared) and (not images_same or drawing_grows):
            return out("uncertain", reasons=["neither page has text; the later page adds to the pictures or the drawing of the earlier page"])
        return None
    if identical:
        if page_identical:
            return out("join", "identical", "duplicate", "high")
        substantial = more_than_a_heading or authored_same
        if unread:
            return out("uncertain", reasons=["the text is identical, but a picture on these pages could not be read, so the pictures were not compared"])
        if not images_kept:  # another picture under the same words: the deck's own number can make it one slide, equal text cannot
            if authored_same:
                return step("same_authored_slide")
            return out("uncertain", reasons=["the text is identical but a picture of the earlier page is replaced or missing"])
        if not compared:
            if substantial:
                return step("identical_text_drawing_not_compared")
            return out("uncertain", reasons=["the text is identical but short, and the drawing of the pages could not be compared"])
        if in_place and drawing_kept and not cover_added and (drawing_grows or not images_same):
            # All that the earlier page shows is on the later page, where it was, and more is shown. With only a heading in common it stays a medium reading.
            relation = "identical_text_drawing_grows" if content_added or not images_same else "identical_text_cover_removed"
            return out("join", relation, "incremental_build", "high" if substantial and (same_title or authored_same) else "medium")
        if substantial and carried:
            return step("identical_text_visual_change")
        if authored_same:
            return step("same_authored_slide")
        if not mostly_same:
            return out("uncertain", reasons=[f"the text is identical but short, and {differs}"])
        return out("uncertain", reasons=["the text is identical but short, and a picture or the drawing changes"])
    if superset:
        if unread:
            return out("uncertain", reasons=["the text grows, but a picture on these pages could not be read, so the pictures were not compared"])
        if not images_kept and not authored_same:
            return out("uncertain", reasons=["the text grows, but a picture of the earlier page is replaced or missing"])
        if compared and not drawing_kept and not carried:
            return out("uncertain", reasons=[f"the text grows, but it is short and {differs}"])
        heading_with_its_figure = same_title and (bool(a.images) or shapes[0] > 0) and images_kept and (drawing_kept or not compared)
        if more_than_a_heading or authored_same or heading_with_its_figure:
            if not images_kept or (compared and not drawing_kept) or cover_added:  # the text grows while something drawn changes or is covered: not everything is known to carry over
                return step("text_superset_visual_change")
            if not compared:  # what the pages draw could not be compared: knowing less never gives the stronger claim
                return step("text_superset_drawing_not_compared")
            firm = ((same_title and more_than_a_heading) or authored_same) and (in_place or extended == 0)
            return out("join", "text_superset", "incremental_build", "high" if firm else "medium")
        return out("uncertain", reasons=["the earlier page shows only the heading that the later page also carries: a section page or a first build"])
    if authored_same:
        if same_title or shrinks or (body_overlap or 0) >= 0.5:
            return step("same_authored_slide")
        return out("uncertain", reasons=["the deck gives both pages the same number or label, but their text is unrelated"])
    if same_words and (same_title or a.body_total >= 3 or (bool(a.images) and images_same)):
        return out("uncertain", reasons=[f"every word of the earlier page recurs on the later page, but not as the same lines in the same order ({kept} of {len(a.line_tokens)} lines are kept): "
                                         "reworded, rearranged or wrapped anew — not established as a build"])
    if shrinks and same_title:
        return out("uncertain", reasons=[f"same heading, and the later page drops text of the earlier page ({removed} of {ta} tokens) without adding any"])
    if same_title and body_overlap is not None and body_overlap >= 0.7 and a.body_total >= 4:
        return out("uncertain", reasons=[f"same heading and most, but not all, of the text carries over: {removed} of {ta} tokens are gone and {added} are new "
                                         "(a changed step, or a build with an edit)"])
    if same_title:
        return out("distinct", reasons=["same heading, different content"])
    if a.images and images_same:
        return out("distinct", reasons=["same picture, different heading and text"])
    return None


def _detect_regions(page: dict) -> list[dict]:
    """Regions of a handout page, only when evident: an upright page that places two or more equal, slide-shaped,
    non-overlapping objects. Anything less clear is left to a correction."""
    size = page.get("size")
    if page["status"] != "ok" or not size or size["height_pt"] <= size["width_pt"] * 1.1:
        return []
    boxes = []
    for placement in page.get("placements", []):
        box = placement["box"]
        if not box:
            continue
        width, height = (box[2] - box[0]) * size["width_pt"], (box[3] - box[1]) * size["height_pt"]
        area = (box[2] - box[0]) * (box[3] - box[1])
        if height > 0 and 1.2 <= width / height <= 1.9 and 0.04 <= area <= 0.45:
            boxes.append((box, width, height))
    if len(boxes) < 2:
        return []
    reference = boxes[0]
    equal = [box for box, width, height in boxes if abs(width - reference[1]) <= 0.03 * reference[1] and abs(height - reference[2]) <= 0.03 * reference[2]]
    if len(equal) < 2 or len(equal) != len(boxes):
        return []
    for i, first in enumerate(equal):
        for second in equal[i + 1:]:
            if min(first[2], second[2]) - max(first[0], second[0]) > 0.005 and min(first[3], second[3]) - max(first[1], second[1]) > 0.005:
                return []
    ordered = sorted(equal, key=lambda box: (round(box[1], 2), box[0]))
    evidence = f"upright page with {len(ordered)} equal slide-shaped objects that do not overlap"
    return [{"region_id": f"r{n}", "order": n, "box": box, "label": None, "title": None, "printed_number": None, "source": "auto", "evidence": evidence} for n, box in enumerate(ordered, start=1)]


def _formula_like(lines: list[str]) -> bool:
    """Whether a line reads like a formula, a relation or a reaction: ``A + B -> C``, ``Km >> [S]``, ``p < 0.05``."""
    for line in lines:
        text = _RE_NOT_FORMULA.sub(" ", _BULLET.sub("", line))
        if any(ch in _FORMULA_MARKS for ch in text) or _RE_FORMULA.search(text):
            return True
    return False


def _role(page: dict, feat: _Feat, first_content_page: int) -> dict:
    """The content role of one page, with the evidence.

    A page that shows a picture (also one written into the page description), any painted shape beyond the template,
    or a line that reads like a formula is never ``blank`` and never ``structural`` on the ground that it has few
    words: it is ``unknown`` or ``instructional`` with a confidence below 0.6 (``uncertain``), and the study flow
    keeps its requirements. Structural, administrative and blank are said only on positive evidence.
    """
    def out(role: str, confidence: float, *evidence: str) -> dict:
        return {"role": role, "confidence": round(confidence, 2), "uncertain": confidence < 0.6, "evidence": list(evidence), "source": "auto",
                "forced_requirements": role not in _NO_FORCED_REQUIREMENT}

    if not feat.ok:
        return out("unknown", 0.0, f"the page could not be inspected ({page.get('reason')})")
    pictures = sum(feat.images.values())
    visual = []
    if pictures:
        visual.append(f"{pictures} picture(s)")
    if feat.shapes:
        visual.append(f"a drawing ({feat.shapes} painted shape(s) beyond what every page carries)")
    elif not feat.marks_known and feat.painted:
        visual.append("a drawing (painted shapes that could not be told apart from the template)")
    words = feat.word_total
    if words == 0:
        if visual:
            return out("unknown", 0.3, "no extractable text, but the page shows " + " and ".join(visual) + "; it is kept as content of unknown role")
        if (page.get("text_operations") or 0) > 0:
            return out("unknown", 0.2, "the page draws text that could not be extracted")
        return out("blank", 0.95 if not page["image_count"] else 0.75, "no text, no picture and no drawing beyond what every page carries")
    title = feat.title or ""
    heading_only = feat.body_total <= 2
    text = "\n".join(feat.lines)
    if words > (_DENSE_UPRIGHT if feat.portrait else _DENSE_SLIDE):
        # Running text (a reading, a handout of prose): a heading word or a lettered list says nothing reliable here.
        return out("instructional", 0.6, f"{words} words of running text; slide roles are not read off a page this dense")
    options = sum(1 for line in feat.lines if _RE_OPTION.match(line))
    equation = _formula_like(feat.lines)
    plain = not visual and not equation  # nothing on the page but its words
    question = _RE_QUESTION.search(title)
    if question or options >= 3:
        if heading_only and plain:
            return out("structural_title_or_section", 0.6, f"heading only: {title!r}")
        why = f"heading {title!r} names an example or question" if question else f"{options} answer options listed"
        return out("worked_example_or_question", 0.8 if question else 0.65, why)
    if _RE_ADMIN.search(title) or _RE_ADMIN_EXACT.match(title.strip()):
        return out("administrative", 0.75, f"heading {title!r} names course administration")
    body_hits = len(_RE_ADMIN_BODY.findall(text))
    if body_hits >= 2:
        if words >= 60:
            return out("mixed", 0.5, f"{body_hits} administrative phrases inside a page of {words} words")
        return out("administrative", 0.6, f"{body_hits} administrative phrases in a short page")
    objectives = _RE_OBJECTIVES.search(title)
    if objectives:
        rest = title[objectives.end():]
        if re.match(r"\s+of\b", rest, re.I) and not _RE_OF_THE_COURSE.match(rest):
            objectives = None  # "Review of probability", "Overview of the immune system": the heading names a topic
    if objectives:
        if heading_only and plain:
            return out("structural_title_or_section", 0.6, f"heading only: {title!r}")
        return out("objectives_or_summary", 0.8, f"heading {title!r} names objectives, an outline or a summary")
    if _RE_CLOSING.match(title.strip()) and words <= 8:
        return out("structural_title_or_section", 0.75, f"closing page: {title!r}")
    if _RE_REFERENCES.search(title):
        return out("administrative", 0.6, f"heading {title!r} names references or credits")
    if words <= 12 or (page["page"] == first_content_page and words <= 30):
        if equation:
            return out("instructional", 0.5, f"sparse text ({words} words) that reads like an equation")
        cover = page["page"] == first_content_page and not any(_BULLET.match(line) for line in feat.lines)
        if visual:
            if cover and _RE_COURSE_HEADER.search(text):
                # Positive evidence of a cover, not the word count; the picture or drawing beside it keeps the reading uncertain.
                return out("structural_title_or_section", 0.55, f"first page with a lecture or course header and {words} words; it also shows " + " and ".join(visual))
            return out("instructional", 0.5, f"short text ({words} words) with " + " and ".join(visual))
        if _RE_SECTION.match(title.strip()):
            return out("structural_title_or_section", 0.8, f"section heading: {title!r}")
        if cover:
            return out("structural_title_or_section", 0.8 if words <= 12 or _RE_COVER.search(text) else 0.55, f"first page: a heading and {words} words, no list, picture, drawing or equation")
        if heading_only:
            # A short heading alone is a divider; a whole sentence alone may be a statement the lecturer talks over.
            return out("structural_title_or_section", 0.65 if words <= 4 else 0.5, f"heading only ({words} words, no picture, drawing or equation): {title!r}")
        if words <= 12:
            return out("instructional", 0.5, f"short text ({words} words) under a heading")
    return out("instructional", 0.75 if words >= 20 else 0.6, f"{words} words of content" + (" with " + " and ".join(visual) if visual else ""))


# ---- assembling the map -------------------------------------------------------------------------------------------
_SOURCE_KEYS = ("source_id", "version_id", "registration", "index_sha256_matches")
_MAX_PRESENTATIONS = 8
_MAX_ALIGNMENT_CELLS = 250_000
_MIN_ALIGNED_PAGES = 3  # fewer matched pages are a coincidence of words, not a sequence
_AMPLE_TEXT = 12  # words after the first line from which equal text is a slide's worth, not a caption's
_TEXT_BASIS = {"source": "embedded_text", "visible_text_established": False,
               "note": "Page text is the text embedded in the file. Reading it cannot tell whether a shape or a picture covers it, so a page may list words it does not "
                       "show. In a group of pages, what is visible on each page was not established: the text of a page says what the file holds there, not what is seen. "
                       "A page counts the fills it paints in the paint of its own background (drawing.fills_in_background_paint): such a fill shows nothing and can only "
                       "hide, so a page with some is where this limitation is most likely to matter. It also counts the text runs it paints in that paint "
                       "(drawing.text_runs_in_background_paint; white counts as the paint of bare paper): words the file holds and the page does not show, as on a build "
                       "that keeps a later line unseen. Both counts are facts about the file, not findings about the screen: white text over a dark picture that does not "
                       "cover the page is counted although it is seen, and text under a picture is not counted although it is hidden."}


def _logical_id(sha256: str, pages: list[int]) -> str:
    return f"ls-{sha256[:10]}-p{pages[0]}" + (f"-{pages[-1]}" if len(pages) > 1 else "")


def _kind_of(filename: str | None) -> str:
    return {".pdf": "pdf", ".pptx": "pptx"}.get(Path(filename or "").suffix.lower(), "other")


def _byte_cap(kind: str, limits: Limits) -> int:
    return limits.max_pptx_bytes if kind == "pptx" else limits.max_bytes


def _empty_base(record: dict, limits: Limits) -> dict[str, Any]:
    return {"contract": SOURCE_MAP_VERSION, "detector_revision": DETECTOR_REVISION, "source": record, "page_index_base": PAGE_INDEX_BASE, "page_count": None,
            "inspection": {"status": "complete", "reason": None, "tool": None, "limits": asdict(limits), "pages_unknown": 0, "transient": False, "attempts": 1, "retry": "on_request"},
            "page_layout": {"kind": "unknown", "landscape_pages": 0, "upright_pages": 0},
            "pdf_labels": {"defined": False, "differ_from_index": False, "repeat": False, "repeat_is_a_signal": False, "distinct": 0},
            "printed_numbers": {"established": False, "numbers_count_slides": False, "evidence": None},
            "pptx": None, "region_box_units": REGION_BOX_UNITS, "text_basis": dict(_TEXT_BASIS),
            "template": {"boilerplate_lines": 0, "template_images": 0, "template_marks": 0, "repeated_in_body": {"images": 0, "marks": 0}},
            "pages": [], "pairs": []}


def _analyse(path: Path, source: dict | None, presentations: list[tuple[Path, dict | None]], limits: Limits, clock: Callable[[], float]) -> dict:
    """The map before any correction: inspection, numbering, alignment, adjacent-page relations, roles."""
    started = clock()
    kind = _kind_of(path.name)
    record: dict[str, Any] = {"source_id": None, "version_id": None, "registration": None, "index_sha256_matches": None, "sha256": None, "filename": path.name, "bytes": None, "kind": kind}
    record.update({key: source[key] for key in _SOURCE_KEYS if source and source.get(key) is not None})
    base = _empty_base(record, limits)

    def stopped(status: str, reason: str, transient: bool = False) -> dict:
        base["inspection"].update(status=status, reason=reason, transient=transient)
        return base

    try:
        record["bytes"] = path.stat().st_size
    except OSError:
        return stopped("unavailable", "missing")
    if record["bytes"] > _byte_cap(kind, limits):
        return stopped("not_inspected", "byte_limit")  # refused on its size alone: such a file is not read, not even for its digest
    try:
        sha = _file_sha256(path)
    except OSError:
        return stopped("unavailable", "read_error", True)
    if source and source.get("sha256") and source["sha256"] != sha:
        raise SourceMapError("identity_mismatch", "the source reference names other bytes than the file that was read")
    record["sha256"] = sha
    record["version_id"] = record["version_id"] or disk_version_id(sha)
    if kind == "pptx":
        inspected = inspect_pptx(path, limits, clock=clock, started=started)
        base["inspection"].update(status="pptx_only" if inspected["status"] == "ok" else inspected["status"], reason=inspected["reason"], tool=inspected["tool"],
                                  transient=inspected["reason"] in _TRANSIENT_REASONS)
        base["pptx"] = {"source": {key: record[key] for key in ("source_id", "version_id", "sha256", "filename")}, "status": inspected["status"], "reason": inspected["reason"],
                        "transient": inspected["reason"] in _TRANSIENT_REASONS,
                        "slide_count": inspected["slide_count"], "hidden_slides": inspected["hidden_slides"], "truncated": inspected["truncated"],
                        "alignment": {"status": "not_attempted", "method": None},
                        "candidates": [], "slides": [dict(slide, matched_pages=[]) for slide in inspected["slides"]]}
        return base
    if kind != "pdf":
        return stopped("not_inspected", "unsupported_type")

    inspected = _inspect_pdf(path, limits, clock, started)
    pages = inspected["pages"]
    base["page_count"] = inspected["page_count"]
    base["inspection"].update(status=inspected["status"], reason=inspected["reason"], tool=inspected["tool"], pages_unknown=sum(1 for p in pages if p["status"] != "ok"),
                              transient=inspected["transient"])
    if not pages:
        return base
    for page in pages:
        page["lines"] = [line.strip() for line in page["text"].splitlines() if line.strip()]
    labels = [page["pdf_label"] for page in pages]
    if inspected["labels_defined"]:
        # A label repeated on consecutive pages names one authored frame only when the labels count: several distinct
        # numbers that never run backwards. One constant label, or a prefix without a number, says nothing about slides.
        numeric = all(re.fullmatch(r"\d{1,6}", label or "") for label in labels)
        counting = numeric and all(int(second) >= int(first) for first, second in zip(labels, labels[1:]))
        repeat = len(set(labels)) < len(labels)
        base["pdf_labels"] = {"defined": True, "differ_from_index": any(label != str(page["page"]) for label, page in zip(labels, pages)), "repeat": repeat,
                              "repeat_is_a_signal": repeat and counting and len(set(labels)) >= 3, "distinct": len(set(labels))}
    printed, evidence = _detect_printed_numbers(pages)
    # Numbers that repeat on consecutive pages count slides (the builds of a slide share its number); numbers that never repeat may count pages.
    by_slide = bool(printed) and (evidence.get("repeats") or 0) > 0
    base["printed_numbers"] = {"established": bool(printed), "numbers_count_slides": by_slide, "evidence": evidence}
    feats, template = _features(pages, printed, evidence.get("margin"))
    base["template"] = template

    upright = sum(1 for feat in feats if feat.ok and feat.portrait)
    landscape = sum(1 for feat in feats if feat.ok) - upright
    measured = upright + landscape
    base["page_layout"] = {"kind": "unknown" if not measured else "slides" if landscape >= 0.8 * measured else "document" if upright >= 0.8 * measured else "mixed",
                           "landscape_pages": landscape, "upright_pages": upright}

    # A presentation of the same unit lends its slide ids only when the alignment supports it. With several
    # presentations in the unit each is tried, and the best supported one is used; every attempt is recorded.
    matches: dict[int, dict] = {}
    tried = []
    for deck_path, deck_source in presentations[:_MAX_PRESENTATIONS]:
        deck = inspect_pptx(deck_path, limits, clock=clock, started=started)
        deck_sha = (deck_source or {}).get("sha256")
        if deck_sha is None and deck["status"] in {"ok", "unreadable"}:  # a file that was refused on its size or for time is not read for a digest either
            try:
                deck_sha = _file_sha256(deck_path)
            except OSError:
                deck_sha = None
        reference = {"source_id": (deck_source or {}).get("source_id"), "version_id": (deck_source or {}).get("version_id") or (disk_version_id(deck_sha) if deck_sha else None),
                     "sha256": deck_sha, "filename": deck_path.name}
        found, alignment = _align_pptx(feats, deck["slides"]) if deck["status"] == "ok" and deck["slides"] else ({}, {"status": "not_attempted", "method": None})
        tried.append((reference, deck, found, alignment))
    if tried:
        supported = [entry for entry in tried if entry[3]["status"] == "supported"]
        reference, deck, matches, alignment = max(supported or tried, key=lambda entry: (len(entry[2]), entry[3].get("candidate_matches") or 0))
        base["pptx"] = {"source": reference, "status": deck["status"], "reason": deck["reason"], "transient": deck["reason"] in _TRANSIENT_REASONS,
                        "slide_count": deck["slide_count"], "hidden_slides": deck["hidden_slides"], "truncated": deck["truncated"], "alignment": alignment,
                        "candidates": [{"source_id": ref["source_id"], "sha256": ref["sha256"], "status": d["status"], "reason": d["reason"], "transient": d["reason"] in _TRANSIENT_REASONS,
                                        "slide_count": d["slide_count"], "alignment": a["status"], "matched_pages": len(f), "used": ref is reference} for ref, d, f, a in tried],
                        "slides": [dict(slide, matched_pages=[page for page, match in sorted(matches.items()) if match["slide_id"] == slide["slide_id"]]) for slide in deck["slides"]]}
        if any(d["reason"] in _TRANSIENT_REASONS for _, d, _, _ in tried):
            # A presentation that could not be read just now (locked, out of time) may lend slide identity on another attempt:
            # the pages are read, but this is not the settled map of the source.
            base["inspection"]["transient"] = True
        for feat in feats:
            feat.pptx_id = matches[feat.page]["slide_id"] if feat.page in matches else None

    labels_signal = base["pdf_labels"]["repeat_is_a_signal"]
    base["pairs"] = [pair for pair in (_pair(a, b, labels_signal, by_slide) for a, b in zip(feats, feats[1:])) if pair is not None]
    first_content = next((feat.page for feat in feats if feat.ok and feat.word_total), 1)
    # A handout shows its grid on most of its pages; equal panels on one page of a reading are a figure, not slides.
    regions = {page["page"]: _detect_regions(page) for page in pages}
    if sum(1 for found in regions.values() if found) < 0.6 * max(1, measured):
        regions = {}
    for page, feat in zip(pages, feats):
        text_identity = hashlib.sha256(_norm(page["text"]).encode("utf-8")).hexdigest()
        unread = sum(1 for image in page["images"] if image["digest_basis"] == "name_and_size")
        base["pages"].append({
            "page": page["page"], "status": page["status"], "reason": page["reason"], "pdf_label": page["pdf_label"],
            "printed_number": printed.get(page["page"]), "printed_number_source": "detected" if page["page"] in printed else None,
            "pptx": matches.get(page["page"]), "title": feat.title if feat.ok else None, "text": page["text"], "text_truncated": page["text_truncated"],
            "word_count": page["word_count"], "content_word_count": feat.word_total, "text_sha256": text_identity,
            # Text and pictures. A picture whose bytes could not be read enters by its name and size only, and the page says so (pictures_unread).
            "content_sha256": _digest([text_identity, sorted(image["sha256"] for image in page["images"])]),
            "images": page["images"], "image_count": page["image_count"], "images_truncated": page["images_truncated"], "pictures_unread": unread,
            "drawing": {"paint_operations": page["paint_operations"], "path_operations": page["path_operations"], "shapes_beyond_template": feat.shapes if feat.ok else None,
                        "fills_in_background_paint": sum(feat.covers.values()) if feat.ok else None, "text_runs_in_background_paint": page["text_in_background_paint"],
                        "marks_compared": feat.marks_known, "visual_sha256": page["visual_sha256"]},
            "size": page["size"], "regions": regions.get(page["page"], []), "retired_regions": [], "role": _role(page, feat, first_content)})
    return base


def _op_gaps(op: dict, page_count: int) -> dict[int, str]:
    """The page boundaries a group or split correction fixes: ``{k: "join" | "cut"}`` for the gap between k and k+1."""
    gaps: dict[int, str] = {}
    if op["type"] == "group":
        first, last = op["pages"][0], op["pages"][-1]
        gaps.update({k: "join" for k in range(first, last)})
        if first > 1:
            gaps[first - 1] = "cut"
        if last < page_count:
            gaps[last] = "cut"
    elif op["type"] == "split":
        if op.get("before_page") is not None:
            gaps[op["before_page"] - 1] = "cut"
        for page in op.get("pages") or []:
            if page > 1:
                gaps[page - 1] = "cut"
            if page < page_count:
                gaps[page] = "cut"
    return gaps


def _op_pages(op: dict) -> list[int] | None:
    """The pages a correction is about (used to say whether it could be re-applied to another source version)."""
    if op["type"] == "group":
        return list(op["pages"])
    if op["type"] == "split":
        pages = set(op.get("pages") or [])
        if op.get("before_page") is not None:
            pages |= {op["before_page"] - 1, op["before_page"]}
        return sorted(pages)
    if op["type"] in {"role", "printed_number", "regions"}:
        return [op["page"]]
    return None


def _validate_op(op: Any, page_count: int | None, known_revisions: list[int], in_force: set[int] | None = None) -> dict:
    """A correction in normal form, or :class:`SourceMapError` (``invalid_override``) saying what is wrong with it."""
    def bad(message: str) -> SourceMapError:
        return SourceMapError("invalid_override", message)

    def page_number(value: Any, what: str = "page") -> int:
        if not isinstance(value, int) or isinstance(value, bool) or not page_count or not 1 <= value <= page_count:
            raise bad(f"{what} must be a page of this source (1 to {page_count})")
        return value

    if not isinstance(op, dict) or op.get("type") not in {"group", "split", "role", "printed_number", "regions", "revert"}:
        raise bad("type must be one of group, split, role, printed_number, regions, revert")
    kind = op["type"]
    if kind == "group":
        pages = op.get("pages")
        if not isinstance(pages, list) or len(pages) < 2:
            raise bad("a group names two or more pages")
        pages = [page_number(p) for p in pages]
        if pages != list(range(pages[0], pages[0] + len(pages))):
            raise bad("a group is a run of adjacent pages in order; pages that are not next to each other are never one slide")
        declared = op.get("kind", "manual_group")
        if declared not in GROUP_KINDS:
            raise bad(f"kind must be one of {', '.join(GROUP_KINDS)}")
        return {"type": "group", "pages": pages, "kind": declared}
    if kind == "split":
        before, pages = op.get("before_page"), op.get("pages")
        if before is None and not pages:
            raise bad("a split names before_page (the page that starts a new slide) or pages (each stands alone)")
        result: dict[str, Any] = {"type": "split"}
        if before is not None:
            if page_number(before, "before_page") < 2:
                raise bad("before_page must be 2 or later: nothing precedes page 1")
            result["before_page"] = before
        if pages:
            if not isinstance(pages, list):
                raise bad("pages must be a list")
            result["pages"] = sorted({page_number(p) for p in pages})
        return result
    if kind == "role":
        if op.get("role") not in ROLES:
            raise bad(f"role must be one of {', '.join(ROLES)}")
        return {"type": "role", "page": page_number(op.get("page")), "role": op["role"]}
    if kind == "printed_number":
        value = op.get("printed_number")
        if value is not None and (not isinstance(value, int) or isinstance(value, bool) or not 0 <= value <= 99999):
            raise bad("printed_number must be a whole number, or null when the page prints none")
        return {"type": "printed_number", "page": page_number(op.get("page")), "printed_number": value}
    if kind == "regions":
        regions = op.get("regions")
        if not isinstance(regions, list) or len(regions) > 24:
            raise bad("regions must be a list of at most 24 boxes (an empty list removes the regions of the page)")
        cleaned = []
        for region in regions:
            box = region.get("box") if isinstance(region, dict) else None
            if (not isinstance(box, list) or len(box) != 4 or any(not isinstance(v, (int, float)) or isinstance(v, bool) or not 0 <= v <= 1 for v in box)
                    or box[0] >= box[2] or box[1] >= box[3]):
                raise bad("a region box is [left, top, right, bottom] as fractions of the page, with left < right and top < bottom")
            number = region.get("printed_number")
            if number is not None and (not isinstance(number, int) or isinstance(number, bool)):
                raise bad("a region's printed_number must be a whole number")
            cleaned.append({"box": [round(float(v), 4) for v in box], "label": str(region["label"])[:120] if region.get("label") else None,
                            "title": str(region["title"])[:200] if region.get("title") else None, "printed_number": number})
        return {"type": "regions", "page": page_number(op.get("page")), "regions": cleaned}
    revisions = op.get("revisions")
    if not isinstance(revisions, list) or not revisions or any(isinstance(r, bool) or r not in known_revisions for r in revisions):
        raise bad("a revert names earlier revisions of this source's corrections")
    if in_force is not None:
        idle = sorted(r for r in set(revisions) if r not in in_force)
        if idle:  # such a record would be stored as applied and change nothing
            raise bad(f"revision {', '.join(str(r) for r in idle)} is not in force (reverted already, or never applied), so there is nothing to revert; "
                      "to bring a reverted correction back, revert the revision that reverted it")
    return {"type": "revert", "revisions": sorted(set(revisions))}


def _author(record: dict) -> tuple[str | None, bool]:
    """Who made a correction — ``learner`` / ``agent`` / unknown — and whether an agent made it on the learner's say."""
    author = record.get("author")
    if isinstance(author, dict):
        return author.get("kind"), bool(author.get("learner_authority"))
    return (author if isinstance(author, str) else None), False


def _by_learner(record: dict) -> bool:
    kind, authority = _author(record)
    return kind == "learner" or (kind == "agent" and authority)


def _in_force(ledger: list[dict], upto: int | None = None) -> set[int]:
    """The revisions whose effect holds. A revert takes its targets out of force; reverting that revert puts them back."""
    force: set[int] = set()
    cancelled: set[int] = set()
    for record in reversed(ledger):  # latest first: whether a revision holds depends only on later reverts that hold
        if (upto is not None and record["revision"] > upto) or record["op"] is None or record["refused"] or record["revision"] in cancelled:
            continue
        force.add(record["revision"])
        if record["op"]["type"] == "revert":
            cancelled.update(record["op"]["revisions"])
    return force


def _timeline(ledger: list[dict]) -> dict[int, dict[int, int]]:
    """After each revision: the revisions whose effect holds, each with the revision at which it last came into force —
    its own number, or that of the revert which put it back. ``0`` is the state before any correction."""
    timeline: dict[int, dict[int, int]] = {0: {}}
    since: dict[int, int] = {}
    for record in ledger:
        since = {revision: since.get(revision, record["revision"]) for revision in _in_force(ledger, record["revision"])}
        timeline[record["revision"]] = since
    return timeline


def _effective(ledger: list[dict], since: dict[int, int]) -> list[dict]:
    """The corrections that hold (reverts aside), in the order in which they came into force. Where two of them speak
    about the same boundary, role, number or regions, the one that came into force later decides: a correction that is
    put back by undoing its undo is the later word, whatever its number."""
    return sorted((record for record in ledger if record["revision"] in since and record["op"]["type"] != "revert"), key=lambda record: (since[record["revision"]], record["revision"]))


def _decided(ledger: list[dict], since: dict[int, int], page_count: int) -> tuple[dict[int, tuple[str, int, str | None]], dict[tuple[str, int], int]]:
    """Which correction decides each page boundary (``{gap: (state, revision, kind of a join)}``) and each role, printed
    number and set of regions of a page (``{(type, page): revision}``)."""
    gaps: dict[int, tuple[str, int, str | None]] = {}
    facets: dict[tuple[str, int], int] = {}
    for record in _effective(ledger, since):
        op = record["op"]
        if op["type"] in {"group", "split"}:
            for gap, state in _op_gaps(op, page_count).items():
                gaps[gap] = (state, record["revision"], op.get("kind") if state == "join" else None)
        else:
            facets[(op["type"], op["page"])] = record["revision"]
    return gaps, facets


def _learner_word(ledger: list[dict], since: dict[int, int], page_count: int) -> tuple[dict[int, int], dict[int, tuple[str, int, str | None]], dict[tuple[str, int], int]]:
    """What the learner's corrections in force settle, the later word replacing the earlier one.

    ``owner``: the corrections that hold on the learner's word — the learner's own, and those a learner's revert put
    back in force — each with the learner's revision it stands on. ``gaps``: page boundaries, with the state the
    learner's word leaves there and the kind the learner gave a join. ``facets``: the role, printed number or regions
    of a page the learner has spoken about.

    A revert is a word too. What it undid is not an agent's to record again: the boundaries of an undone group or
    split are held at the other state, and the role, number or regions of that page are the learner's. What it put
    back in force stands like a correction of the learner's own.
    """
    by_revision = {record["revision"]: record for record in ledger}
    owner: dict[int, int] = {}
    gaps: dict[int, tuple[str, int, str | None]] = {}
    facets: dict[tuple[str, int], int] = {}

    def speak(record: dict, by: int, depth: int = 0) -> None:
        owner[record["revision"]] = by
        op = record["op"]
        if op["type"] in {"group", "split"}:
            for gap, state in _op_gaps(op, page_count).items():
                gaps[gap] = (state, by, op.get("kind") if state == "join" else None)
        elif op["type"] != "revert":
            facets[(op["type"], op["page"])] = by
        elif depth < 64:
            for target in op["revisions"]:
                undone = by_revision.get(target)
                if undone is None or undone["op"] is None or undone["refused"]:
                    continue
                other = undone["op"]
                if other["type"] == "revert":  # an undo is undone: what it had taken out of force holds again, now on the learner's word
                    for revision in other["revisions"]:
                        if revision in since:
                            speak(by_revision[revision], by, depth + 1)
                elif other["type"] in {"group", "split"}:
                    for gap, state in _op_gaps(other, page_count).items():
                        gaps[gap] = ("cut" if state == "join" else "join", by, None)
                else:
                    facets[(other["type"], other["page"])] = by

    for record in sorted((record for record in ledger if record["revision"] in since and _by_learner(record)), key=lambda record: (since[record["revision"]], record["revision"])):
        speak(record, record["revision"])
    return owner, gaps, facets


def _learner_conflicts(op: dict, ledger: list[dict], page_count: int) -> list[int]:
    """The learner's corrections in force that this correction would undo, contradict or take the place of (their
    revision numbers). ``ledger`` is every earlier record; the correction is measured by what recording it would do:
    nothing the learner's word holds may go out of force, no boundary the learner settled may be decided otherwise
    (nor a join be given another kind), and no role, number or regions the learner spoke about may get another
    deciding correction. That covers the correction itself, a revert of a learner's correction, a revert that would
    put back what the learner undid, and the same correction recorded again after the learner's undo."""
    before = _timeline(ledger)[ledger[-1]["revision"] if ledger else 0]
    owner, gaps, facets = _learner_word(ledger, before, page_count)
    if not owner:
        return []
    revision = (ledger[-1]["revision"] if ledger else 0) + 1
    after_ledger = [*ledger, {"revision": revision, "op": op, "refused": False}]
    after = _timeline(after_ledger)[revision]
    against = {by for held, by in owner.items() if held not in after}
    gaps_before, facets_before = _decided(ledger, before, page_count)
    gaps_after, facets_after = _decided(after_ledger, after, page_count)
    for gap, (state, by, kind) in gaps.items():
        decided = gaps_after.get(gap)
        if decided is not None and decided != gaps_before.get(gap) and (decided[0] != state or (kind is not None and decided[2] != kind)):
            against.add(by)
    for facet, by in facets.items():
        if facets_after.get(facet) != facets_before.get(facet):
            against.add(by)
    return sorted(against)


def _refusal(against: list[int]) -> str:
    return (f"the learner's correction r{', r'.join(str(r) for r in against)} is in force here; a correction made by an agent never replaces or undoes one made by the learner, "
            "and never records again what the learner undid. It is recorded only when the learner makes it, or when the call says it is made on the learner's authority")


def _ledger(records: list[dict], page_count: int) -> tuple[list[dict], list[dict]]:
    """The correction records in order and in normal form, each with what became of it: ``op`` is ``None`` for a record
    that is not valid, and ``refused`` is set for a record that is not the learner's — an agent's, or one that does not
    say who made it — and goes against a learner's correction in force (neither has any effect). Also returns the
    problems found."""
    ledger: list[dict] = []
    problems: list[dict] = []
    known: list[int] = []
    for position, record in enumerate(records, start=1):
        record = {**record, "revision": record.get("revision", position)}  # records given without numbers are numbered in order
        force = _in_force(ledger)
        refused = False
        try:
            op = _validate_op(record.get("op"), page_count, known, force)
        except SourceMapError as exc:
            problems.append({"code": "invalid_override", "revision": record["revision"], "message": exc.message})
            op = None
        if op is not None and not _by_learner(record):
            against = _learner_conflicts(op, ledger, page_count)
            if against:
                refused = True
                problems.append({"code": "learner_correction_in_force", "revision": record["revision"], "message": _refusal(against)})
        known.append(record["revision"])
        ledger.append({**record, "op": op, "refused": refused})
    return ledger, problems


def _replaced_by(op: dict, live: list[dict], page_count: int, agent_alone: bool) -> list[int]:
    """Earlier corrections in force whose whole effect the new one takes over. An agent's correction never takes over a learner's."""
    if op["type"] == "revert":
        return list(op["revisions"])
    replaced = []
    gaps = _op_gaps(op, page_count)
    for record in live:
        other = record["op"]
        if other["type"] == "revert" or (agent_alone and _by_learner(record)):
            continue
        if op["type"] in {"role", "printed_number", "regions"}:
            if other["type"] == op["type"] and other["page"] == op["page"]:
                replaced.append(record["revision"])
        elif other["type"] in {"group", "split"}:
            other_gaps = _op_gaps(other, page_count)
            if other_gaps and set(other_gaps) <= set(gaps):
                replaced.append(record["revision"])
    return replaced


def _runs(page_count: int, joined: set[int]) -> list[list[int]]:
    runs, current = [], [1]
    for page in range(2, page_count + 1):
        if page - 1 in joined:
            current.append(page)
        else:
            runs.append(current)
            current = [page]
    runs.append(current)
    return runs


def _same_box(first: list[float], second: list[float]) -> bool:
    return all(abs(a - b) <= 0.002 for a, b in zip(first, second))


def _overlap(first: list[float], second: list[float]) -> float:
    """Intersection over union of two boxes."""
    width, height = min(first[2], second[2]) - max(first[0], second[0]), min(first[3], second[3]) - max(first[1], second[1])
    if width <= 0 or height <= 0:
        return 0.0
    shared = width * height
    return shared / ((first[2] - first[0]) * (first[3] - first[1]) + (second[2] - second[0]) * (second[3] - second[1]) - shared)


def _regions_of(page: dict, ledger: list[dict], timeline: dict[int, dict[int, int]]) -> tuple[list[dict], list[dict]]:
    """The regions of a page after its corrections, and the regions that are gone.

    A region id keeps naming the box it named: a correction that lists the same box again keeps its id whatever the
    order, a new box gets an id that was never used on this page, and an id that is no longer listed is retired with
    the box it had — so an old address still answers, and says where the region was and which region now covers it.
    """
    number = page["page"]
    relevant = [record for record in ledger if record["op"] is not None and not record["refused"] and record["op"]["type"] in {"regions", "revert"}]
    automatic = [dict(region) for region in page["regions"]]
    if not any(record["op"]["type"] == "regions" and record["op"]["page"] == number for record in relevant):
        return automatic, []
    issued: dict[str, list[float]] = {region["region_id"]: region["box"] for region in automatic}  # every id ever given on this page, with its last box
    retired: dict[str, int] = {}
    current, governing = automatic, None
    for record in relevant:
        holder = next((r for r in reversed(_effective(ledger, timeline[record["revision"]])) if r["op"]["type"] == "regions" and r["op"]["page"] == number), None)
        if (holder["revision"] if holder else None) == governing:
            continue
        governing = holder["revision"] if holder else None
        if holder is None:
            current = automatic
        else:
            current, used = [], set()
            for order, region in enumerate(holder["op"]["regions"], start=1):
                region_id = next((key for key, box in issued.items() if key not in used and _same_box(box, region["box"])), None)
                if region_id is None:
                    region_id = f"r{max((int(key[1:]) for key in issued), default=0) + 1}"
                issued[region_id] = region["box"]
                used.add(region_id)
                current.append({"region_id": region_id, "order": order, "box": region["box"], "label": region["label"], "title": region["title"],
                                "printed_number": region["printed_number"], "source": "override", "evidence": f"correction r{holder['revision']}"})
        now = {region["region_id"] for region in current}
        for region_id in issued:
            if region_id in now:
                retired.pop(region_id, None)
            else:
                retired.setdefault(region_id, record["revision"])
    gone = []
    for region_id, revision in sorted(retired.items(), key=lambda item: int(item[0][1:])):
        covering = max(current, key=lambda region: _overlap(region["box"], issued[region_id]), default=None)
        gone.append({"region_id": region_id, "box": issued[region_id], "retired_by_revision": revision,
                     "resolves_to": covering["region_id"] if covering is not None and _overlap(covering["box"], issued[region_id]) >= 0.5 else None})
    return current, gone


def _apply(base: dict, records: list[dict], problems: list[dict] | None = None) -> dict:
    """The map at the revision that ``records`` reach. ``base`` is the uncorrected analysis; it is not modified."""
    result = copy.deepcopy(base)
    sha = result["source"]["sha256"] or ""  # a file that was not read has no pages, so no id is ever made from this
    pages = result["pages"]
    count = len(pages)
    ledger, found = _ledger(records, count)
    problems = [*(problems or []), *found]
    timeline = _timeline(ledger)

    def effective(upto: int) -> list[dict]:  # in the order in which the corrections came into force: the later word decides
        return _effective(ledger, timeline[upto])

    auto = {pair["pages"][0]: pair for pair in result["pairs"] if pair["decision"] == "join"}

    def forced_gaps(upto: int) -> dict[int, tuple[str, int, str | None]]:
        forced: dict[int, tuple[str, int, str | None]] = {}
        for record in effective(upto):
            for gap, state in _op_gaps(record["op"], count).items():
                forced[gap] = (state, record["revision"], record["op"].get("kind"))
        return forced

    def grouping(upto: int) -> list[list[int]]:
        forced = forced_gaps(upto)
        joined = {gap for gap in range(1, count) if (forced[gap][0] == "join" if gap in forced else gap in auto)}
        return _runs(count, joined) if count else []

    head = ledger[-1]["revision"] if ledger else 0
    forced = forced_gaps(head)
    live = effective(head)
    role_ops = {r["op"]["page"]: r for r in live if r["op"]["type"] == "role"}
    number_ops = {r["op"]["page"]: r for r in live if r["op"]["type"] == "printed_number"}
    for page in pages:
        number = page["page"]
        if number in role_ops:
            automatic = page["role"]
            role = role_ops[number]["op"]["role"]
            page["role"] = {"role": role, "confidence": 1.0, "uncertain": False, "evidence": [f"correction r{role_ops[number]['revision']}: {role_ops[number].get('reason') or ''}".strip()],
                            "source": "override", "forced_requirements": role not in _NO_FORCED_REQUIREMENT, "override_revision": role_ops[number]["revision"], "automatic": automatic}
        if number in number_ops:
            page["printed_number_detected"] = page["printed_number"]
            page["printed_number"] = number_ops[number]["op"]["printed_number"]
            page["printed_number_source"] = "override"
        page["regions"], page["retired_regions"] = _regions_of(page, ledger, timeline)

    uncertain: dict[int, list[tuple[int, list[str]]]] = {}
    for pair in result["pairs"]:
        first, second = pair["pages"]
        if pair["decision"] == "uncertain" and first not in forced:
            uncertain.setdefault(first, []).append((second, pair["reasons"]))
            uncertain.setdefault(second, []).append((first, pair["reasons"]))
    pair_at = {pair["pages"][0]: pair for pair in result["pairs"]}
    titles_seen: dict[str, list[int]] = {}
    slides = []
    for ordinal, run in enumerate(grouping(head), start=1):
        inner = run[:-1]
        edges = [gap for gap in (run[0] - 1, run[-1]) if 1 <= gap < count]
        join_revisions = sorted({forced[gap][1] for gap in inner if gap in forced and forced[gap][0] == "join"})
        # A forced boundary counts as a correction of this slide when the automatic reading would have joined there or was unsure.
        cut_revisions = sorted({forced[gap][1] for gap in edges if gap in forced and forced[gap][0] == "cut"
                                and (gap in auto or (gap in pair_at and pair_at[gap]["decision"] == "uncertain"))})
        evidence = []
        for gap in inner:
            if gap in auto:
                pair = auto[gap]
                evidence.append({"pages": pair["pages"], "relation": pair["relation"], "kind": pair["kind"], "signals": pair["signals"], "confidence": pair["confidence"],
                                 "measures": pair["measures"], "notes": pair["notes"]})
            else:
                note = pair_at[gap]["reasons"] if gap in pair_at else []
                evidence.append({"pages": [gap, gap + 1], "relation": "override", "kind": None, "signals": ["override"], "confidence": None, "measures": None,
                                 "notes": [f"joined by correction r{forced[gap][1]}", *(f"automatic reading: {reason}" for reason in note)]})
        reasons: list[str] = []
        neighbours: list[int] = []
        cumulative: bool | None = None  # whether the last page shows everything of the pages before it
        if len(run) == 1:
            kind = "single"
            for other, why in uncertain.get(run[0], []):
                neighbours.append(other)
                reasons.extend(reason for reason in why if reason not in reasons)
            state = "override" if cut_revisions else "uncertain" if reasons else "none"
            confidence = None
        elif join_revisions:
            declared = [forced[gap][2] for gap in inner if gap in forced and forced[gap][0] == "join" and forced[gap][2]]
            kind, state, confidence = (declared[-1] if declared else "manual_group"), "override", None
        else:
            kinds = {auto[gap]["kind"] for gap in inner}
            # One changed state among the pages makes the whole group a sequence of states: it is then not a build and not a copy.
            kind = "animation_steps" if "animation_steps" in kinds else "incremental_build" if "incremental_build" in kinds else "duplicate"
            state = "inferred"
            confidence = "high" if kind != "animation_steps" and all(auto[gap]["confidence"] == "high" for gap in inner) else "medium"
            cumulative = kind != "animation_steps"
        title = next((pages[n - 1]["title"] for n in reversed(run) if pages[n - 1]["title"]), None)
        representative = pages[run[-1] - 1]
        key = _norm(title or "")
        earlier = [n for n in titles_seen.get(key, []) if n < run[0] - 1] if key else []
        if key:
            titles_seen.setdefault(key, []).append(run[-1])
        slide_ids: list[dict] = []
        for n in run:
            match = pages[n - 1]["pptx"]
            if match and {"slide_id": match["slide_id"], "ordinal": match["ordinal"]} not in slide_ids:
                slide_ids.append({"slide_id": match["slide_id"], "ordinal": match["ordinal"]})
        slide = {"logical_slide_id": _logical_id(sha, run), "ordinal": ordinal, "kind": kind, "grouping": state, "confidence": confidence, "pages": run,
                 "representative_page": run[-1], "cumulative": cumulative, "text_visibility": "not_established", "title": title,
                 "printed_numbers": sorted({pages[n - 1]["printed_number"] for n in run if pages[n - 1]["printed_number"] is not None}),
                 "pptx_slides": slide_ids, "role": representative["role"]["role"], "role_uncertain": representative["role"]["uncertain"], "evidence": evidence, "uncertain_reasons": reasons, "uncertain_with": sorted(set(neighbours)),
                 "repeats_title_of": earlier[:20], "overrides": sorted(set(join_revisions) | set(cut_revisions))}
        slides.append(slide)
        for index, n in enumerate(run, start=1):
            pages[n - 1].update(logical_slide_id=slide["logical_slide_id"], build_index=index, build_count=len(run))
    result["logical_slides"] = slides

    # Every id a revision ever issued stays resolvable: an id that is no longer current becomes an alias.
    current = {slide["logical_slide_id"]: slide for slide in slides}
    by_page = {n: slide["logical_slide_id"] for slide in slides for n in slide["pages"]}
    retired: dict[str, tuple[list[int], int]] = {}
    previous = {_logical_id(sha, run): run for run in grouping(0)}
    for record in ledger:
        now = {_logical_id(sha, run): run for run in grouping(record["revision"])}
        for logical_id, run in previous.items():
            if logical_id not in now:
                retired[logical_id] = (run, record["revision"])
        previous = now
    aliases = []
    for logical_id, (run, revision) in sorted(retired.items(), key=lambda item: (item[1][0][0], item[1][0][-1])):
        if logical_id in current:
            continue
        targets = list(dict.fromkeys(by_page[n] for n in run))
        aliases.append({"alias": logical_id, "pages": run, "retired_by_revision": revision, "resolves_to": targets, "primary": targets[0]})
    result["aliases"] = aliases

    force = _in_force(ledger)
    by_revision = {record["revision"]: record for record in ledger}
    reverted_by = {target: record["revision"] for record in ledger if record["revision"] in force and record["op"]["type"] == "revert" for target in record["op"]["revisions"]}
    superseded = {x for record in ledger if record["revision"] in force and record["op"]["type"] != "revert" for x in (record.get("replaces") or [])}
    history = []
    for record in ledger:
        op, revision = record["op"], record["revision"]
        status = ("invalid" if op is None else "refused" if record["refused"] else "reverted" if revision not in force
                  else "superseded" if revision in superseded else "applied")
        entry = {"revision": revision, "type": op["type"] if op else None, "op": op, "author": record.get("author"), "reason": record.get("reason"),
                 "created_at": record.get("created_at"), "supersedes": record.get("supersedes"), "replaces": record.get("replaces") or [], "status": status,
                 "in_force": revision in force, "reverted_by": reverted_by.get(revision)}
        if op is not None and op["type"] == "revert":
            # Reverting a revert puts back what that revert had undone; the history names it.
            entry["restores"] = sorted({target for undone in op["revisions"] if (by_revision[undone]["op"] or {}).get("type") == "revert"
                                        for target in by_revision[undone]["op"]["revisions"]} & force) if revision in force else []
        history.append(entry)
    result["overrides"] = {"contract": OVERRIDE_CONTRACT, "revision": head, "history": history, "problems": problems}
    result.setdefault("other_source_versions", [])
    result["map_revision"] = _digest({key: result[key] for key in ("contract", "detector_revision", "pages", "logical_slides", "aliases", "pptx", "pdf_labels", "printed_numbers", "page_layout")}
                                     | {"sha256": result["source"]["sha256"], "overrides_revision": head})[:16]
    return _plain(result)


def build_source_map(path: Path, *, source: dict | None = None, pptx_path: Path | None = None, pptx_source: dict | None = None,
                     pptx_candidates: list[tuple[Path, dict | None]] | None = None, overrides: list[dict] | None = None,
                     limits: Limits | None = None, clock: Callable[[], float] | None = None) -> dict[str, Any]:
    """Inspect one source file and return its map. Pure: reads the file(s), writes nothing.

    ``source`` may carry the ``source_id`` / ``version_id`` the caller knows; the sha256 is always taken from the
    bytes that are read (a file over the byte bound is refused before it is read, so its map has no sha256).
    ``pptx_path`` names a presentation of the same unit (``pptx_candidates`` several, as ``(path, source)`` pairs):
    slide ids are attached only where a text/title sequence alignment of at least three pages supports them, and with
    several candidates the best supported one is used. ``overrides`` are correction records (as stored, or the
    ``overrides.history`` of an earlier map) applied in order; an agent's record that contradicts a learner's record
    in force is reported in ``overrides.problems`` and not applied. ``inspection.transient`` says that the inspection
    was cut short by something of the moment (a read error, the time bound) and another call may give more.
    """
    presentations = [(Path(deck), deck_source) for deck, deck_source in pptx_candidates or []]
    if pptx_path is not None:
        presentations.insert(0, (Path(pptx_path), pptx_source))
    base = _analyse(Path(path), source, presentations, limits or Limits(), clock or time.monotonic)
    return _apply(base, [dict(record) for record in overrides or []])


# ---- stored corrections and the cache (the app's own area only) ------------------------------------------------------
def _source_dir(settings: Settings, course_id: str, source_id: str) -> Path:
    name = source_id if _SAFE_ID.match(source_id or "") else "id-" + hashlib.sha256((source_id or "").encode("utf-8")).hexdigest()[:20]
    return settings.app_area(course_id) / "source-map" / name


def _load_chain(directory: Path, sha256: str | None) -> tuple[list[dict], list[str], list[dict]]:
    """The verified run of correction revisions in a folder: records, the digest of each file, and what was wrong.

    A revision must follow its predecessor by number and name the predecessor's bytes; reading stops at the first
    one that does not, so a damaged tail is reported instead of being applied.
    """
    records: list[dict] = []
    digests: list[str] = []
    problems: list[dict] = []
    revision = 1
    while True:
        path = directory / f"r{revision:04d}.json"
        if not path.is_file():
            break
        data = path.read_bytes()
        try:
            record = c.parse_json(data)
        except c.ContractError:
            problems.append({"code": "unreadable_revision", "revision": revision})
            break
        expected = None if revision == 1 else {"revision": revision - 1, "sha256": digests[-1]}
        if (not isinstance(record, dict) or record.get("contract") != OVERRIDE_CONTRACT or record.get("revision") != revision or record.get("supersedes") != expected
                or (sha256 is not None and (record.get("source") or {}).get("sha256") != sha256)):
            problems.append({"code": "broken_revision_chain", "revision": revision})
            break
        records.append(record)
        digests.append(c.sha256_hex(data))
        revision += 1
    return records, digests, problems


def _content_digest(stored: dict) -> str:
    """The digest a cached map carries of its own content (everything but the digest itself)."""
    cache = {key: value for key, value in (stored.get("cache") or {}).items() if key != "content_sha256"}
    return _digest({**{key: value for key, value in stored.items() if key != "cache"}, "cache": cache})


def _read_cached(path: Path) -> dict | None:
    """A cached map, or ``None`` when the file is missing, not JSON, or no longer what was written (its own digest does not match)."""
    try:
        stored = json.loads(path.read_text(encoding="utf-8"))
        if isinstance(stored, dict) and isinstance(stored.get("cache"), dict) and stored["cache"].get("content_sha256") == _content_digest(stored):
            return stored
    except (OSError, ValueError, TypeError):
        pass
    return None


def _other_versions(source_dir: Path, current: str, pages: list[dict]) -> tuple[list[dict], list[Any]]:
    """Maps and corrections kept for other bytes of the same source: they do not apply to the current file."""
    reports: list[dict] = []
    fingerprint: list[Any] = []
    if not source_dir.is_dir():
        return reports, fingerprint
    now = {page["page"]: page["content_sha256"] for page in pages if page["status"] == "ok" and not page.get("pictures_unread")}
    for directory in sorted(p for p in source_dir.iterdir() if p.is_dir() and _VERSION_DIR.match(p.name) and p.name != current):
        maps = sorted((p for p in directory.iterdir() if _MAP_FILE.match(p.name)), key=lambda p: (int(_MAP_FILE.match(p.name).group(1)), p.stat().st_mtime_ns))
        records, digests, _ = _load_chain(directory / "overrides", None)
        fingerprint.append([directory.name, [p.name for p in maps], digests])
        old: dict = {}
        if maps:
            try:
                old = json.loads(maps[-1].read_text(encoding="utf-8"))
            except (OSError, ValueError):
                old = {}
        before = {page["page"]: page.get("content_sha256") for page in old.get("pages") or [] if page.get("status") == "ok" and not page.get("pictures_unread")}
        not_applied = []
        for record in records:
            op = record.get("op") or {}
            touched = _op_pages(op) if op.get("type") else None
            same = None if touched is None or not before else all(n in before and before[n] == now.get(n) for n in touched)
            not_applied.append({"revision": record["revision"], "type": op.get("type"), "pages": touched, "author": record.get("author"), "reason": record.get("reason"),
                                "status": "not_applied_source_replaced", "same_content_at_these_pages": same})
        ids = [slide["logical_slide_id"] for slide in old.get("logical_slides") or []] + [alias["alias"] for alias in old.get("aliases") or []]
        reports.append({"sha256": (old.get("source") or {}).get("sha256") or next(((r.get("source") or {}).get("sha256") for r in records), None), "sha256_prefix": directory.name,
                        "version_id": (old.get("source") or {}).get("version_id"), "map_revision": old.get("map_revision"), "page_count": old.get("page_count"),
                        "logical_slide_ids_not_applicable": sorted(set(ids)), "overrides_not_applied": not_applied,
                        "note": "Ids and corrections of other bytes of this source. None is applied to the current file; a correction marked same_content_at_these_pages can be made again."})
    return reports, fingerprint


def _uncached(item: dict, limits: Limits, status: str, reason: str, *, size: int | None = None, transient: bool = False) -> dict:
    """The map of a file that was not read (missing, empty, over the byte bound, or not readable just now). Nothing is kept
    for such a state, so the next call looks at the file again."""
    record = {"source_id": item.get("source_id"), "version_id": item.get("version_id"), "registration": item.get("registration"), "index_sha256_matches": None,
              "sha256": None, "filename": item.get("filename"), "bytes": size if size is not None else item.get("bytes"), "kind": _kind_of(item.get("filename"))}
    base = _empty_base(record, limits)
    base["inspection"].update(status=status, reason=reason, transient=transient, retry="automatic")
    result = _apply(base, [])
    result["cache"] = None
    return result


def _companions(root: Path, item: dict, lecture: dict | None, companions: list[dict] | None) -> list[dict]:
    """The readable ``.pptx`` decks of the same unit (or the candidates the caller names), at most eight."""
    if Path(item.get("filename") or "").suffix.lower() != ".pdf":
        return []
    pool = companions if companions is not None else (lecture or {}).get("decks") or []
    found = [deck for deck in pool if deck.get("source_id") != item.get("source_id") and deck.get("path") and Path(deck.get("filename") or "").suffix.lower() == ".pptx"
             and (root / deck["path"]).is_file() and (root / deck["path"]).stat().st_size > 0]
    return sorted(found, key=lambda deck: deck["source_id"])[:_MAX_PRESENTATIONS]


def load_or_build(settings: Settings, course_id: str, item: dict, *, lecture: dict | None = None, companions: list[dict] | None = None,
                  limits: Limits | None = None, refresh: bool = False, clock: Callable[[], float] | None = None) -> dict[str, Any]:
    """The current map of an inventory item, from the cache when its inputs are unchanged.

    ``item`` is a source of :func:`lecture_brain.ownerio.inventory.build_inventory`. Pass the item's ``lecture`` (or
    ``companions``) so that the presentations of the same unit are tried for slide identity. The cache and the
    correction revisions live under ``settings.app_area(course_id) / "source-map" / <source id> / <sha256 prefix>``;
    nothing is written into the course folder. A file with other bytes gets its own folder and a new map whose
    ``other_source_versions`` lists the ids and corrections that no longer apply.

    What is kept, and for how long:

    * A file over the byte bound, missing or empty is not read (not even hashed) and nothing is kept for it.
    * A file that cannot be opened just now (locked, a read error) gives ``inspection.status = "unavailable"`` with
      ``transient``; nothing is kept and an earlier good map is left as it is.
    * An inspection cut short by something of the moment — a read error on a page, the time bound, a presentation of
      the unit that could not be opened or was not reached in time — gives a map marked ``inspection.transient``
      (``retry = "automatic"``, ``cache`` is ``None``). It is never stored, however often it happens: every call
      inspects the file again until it reads through, and a settled map stored earlier is left as it is. The
      corrections of the source are applied to such a map as to any other. ``inspection.attempts`` counts the
      inspections in a row that ended this way in this process, so that a caller can decide how often to ask: a file
      that always runs out of time costs ``limits.max_seconds`` on every call.
    * ``refresh=True`` always inspects again and replaces what is cached — unless that inspection is itself cut short.
    * A cached file is used only when the digest it carries of its own content still matches; otherwise it is rebuilt.
      A cache that cannot be written does not fail the call.

    A first build can take long for a large file (tens of seconds for several hundred pages): callers on a request
    path should have maps built by the preparation queue.
    """
    limits = limits or Limits()
    root = settings.course(course_id).root
    if not item.get("path"):
        return _uncached(item, limits, "unavailable", "no_file")
    path = root / item["path"]
    try:
        if not path.is_file():
            return _uncached(item, limits, "unavailable", "missing")
        size = path.stat().st_size
        if size == 0:
            return _uncached(item, limits, "unavailable", "empty")
        if size > _byte_cap(_kind_of(path.name), limits):
            return _uncached(item, limits, "not_inspected", "byte_limit", size=size)  # before the file is hashed
        sha = _file_sha256(path)
    except OSError:
        return _uncached(item, limits, "unavailable", "read_error", transient=True)
    declared = item.get("sha256")
    source = {"source_id": item["source_id"], "version_id": item["version_id"] if declared == sha and item.get("version_id") else disk_version_id(sha),
              "registration": item.get("registration"), "index_sha256_matches": (declared == sha) if declared else None, "sha256": sha}
    presentations: list[tuple[Path, dict | None]] = []
    partners: list[str] = []
    partner_unread = False
    for deck in _companions(root, item, lecture, companions):
        deck_path = root / deck["path"]
        try:
            deck_size = deck_path.stat().st_size
            if deck_size > limits.max_pptx_bytes:  # refused on its size: named by that, never read for a digest
                presentations.append((deck_path, {"source_id": deck["source_id"], "sha256": None, "version_id": None}))
                partners.append(f"oversized:{deck['source_id']}:{deck_size}")
                continue
            deck_sha = _file_sha256(deck_path)
        except OSError:  # the presentation is there but cannot be read just now: it is still tried, and whatever comes of it is not kept
            presentations.append((deck_path, {"source_id": deck["source_id"], "sha256": None, "version_id": None}))
            partners.append(f"unread:{deck['source_id']}")
            partner_unread = True
            continue
        presentations.append((deck_path, {"source_id": deck["source_id"], "sha256": deck_sha,
                                          "version_id": deck["version_id"] if deck.get("sha256") == deck_sha and deck.get("version_id") else disk_version_id(deck_sha)}))
        partners.append(deck_sha)
    partners.sort()

    source_dir = _source_dir(settings, course_id, item["source_id"])
    version_dir = source_dir / sha[:16]
    records, digests, problems = _load_chain(version_dir / "overrides", sha)
    suffix = f"-c{_digest(partners)[:8]}" if partners else ""  # a map built with other presentations is another cache entry
    base_digest = _digest({"contract": SOURCE_MAP_VERSION, "detector": DETECTOR_REVISION, "source": source, "limits": asdict(limits), "presentations": partners})

    def read(revision: int) -> dict | None:
        return None if refresh else _read_cached(version_dir / f"map-o{revision:04d}{suffix}.json")

    def base_of(cached: dict) -> dict:
        """The uncorrected analysis inside a cached revision-0 map."""
        base = copy.deepcopy({key: value for key, value in cached.items() if key not in {"logical_slides", "aliases", "overrides", "other_source_versions", "map_revision", "cache"}})
        for page in base["pages"]:
            for key in ("logical_slide_id", "build_index", "build_count"):
                page.pop(key, None)
        return base

    zero = read(0)
    if zero is not None and not (zero["cache"].get("base_digest") == base_digest and zero["cache"].get("analysis_digest") and (zero.get("overrides") or {}).get("revision") == 0
                                 and not zero["inspection"].get("transient")):  # a map cut short is never the stored map, whoever wrote the file
        zero = None
    if zero is not None:
        base = base_of(zero)
        analysis_digest = zero["cache"]["analysis_digest"]
    else:
        base = _analyse(path, source, presentations, limits, clock or time.monotonic)
        inspection = base["inspection"]
        if partner_unread:
            inspection["transient"] = True
        unread = base["source"]["sha256"] is None or inspection["status"] == "unavailable"
        attempt_key = (str(version_dir), suffix, base_digest)
        inspection["attempts"] = _PASSING_FAILURES.get(attempt_key, 0) + 1
        if unread or inspection["transient"]:
            # Something of the moment (the file or a presentation could not be read, the time ran out): this is not the map of the
            # source. Nothing is written, an earlier settled map is left as it is, and the next call inspects again.
            _PASSING_FAILURES[attempt_key] = inspection["attempts"]
            while len(_PASSING_FAILURES) > 512:
                _PASSING_FAILURES.pop(next(iter(_PASSING_FAILURES)))
            inspection["retry"] = "automatic"
            result = _apply(base, [] if unread else records, None if unread else problems)
            if not unread:  # corrections belong to the bytes, so they are shown; the ids and corrections of other bytes are listed as always
                result["other_source_versions"] = _other_versions(source_dir, version_dir.name, result["pages"])[0]
            result["cache"] = None
            return _plain(result)
        _PASSING_FAILURES.pop(attempt_key, None)
        inspection["attempts"], inspection["retry"] = 1, "on_request"  # the settled map: what is stored does not depend on what went before
        base = _plain(base)
        analysis_digest = _digest(base)
    others, fingerprint = _other_versions(source_dir, version_dir.name, base["pages"])

    def finish(revision: int) -> dict:
        inputs_digest = _digest({"base": base_digest, "analysis": analysis_digest, "overrides": digests[:revision], "others": fingerprint})
        cached = zero if revision == 0 and zero is not None else read(revision)
        if cached is not None and cached["cache"].get("inputs_digest") == inputs_digest:
            return cached
        result = _apply(base, records[:revision], problems if revision == len(records) else None)
        result["other_source_versions"] = others
        result["cache"] = {"key": {"sha256": sha, "map_version": SOURCE_MAP_VERSION, "overrides_revision": revision}, "inputs_digest": inputs_digest, "base_digest": base_digest,
                           "analysis_digest": analysis_digest}
        result = _plain(result)
        result["cache"]["content_sha256"] = _content_digest(result)
        try:
            replace_atomic(version_dir / f"map-o{revision:04d}{suffix}.json", c.dump_json(result))
        except OSError:
            pass  # a cache that cannot be written costs the next call an inspection, not this call its answer
        return result

    if records and (zero is None or zero["cache"].get("inputs_digest") != _digest({"base": base_digest, "analysis": analysis_digest, "overrides": [], "others": fingerprint})):
        finish(0)  # keep the uncorrected map beside the corrected ones: the history of the mapping stays readable
    return finish(len(records))


def apply_override(settings: Settings, course_id: str, item: dict, op: dict, *, author: str, reason: str, author_ref: str | None = None, learner_authority: bool = False,
                   lecture: dict | None = None, companions: list[dict] | None = None, now: str | None = None, limits: Limits | None = None) -> dict[str, Any]:
    """Record one correction as the next immutable revision for this source version and return the new map.

    ``op`` is one of::

        {"type": "group", "pages": [7, 8, 9], "kind": "manual_group" | "incremental_build" | "duplicate" | "animation_steps"}
        {"type": "split", "before_page": 9}            # page 9 starts a new logical slide
        {"type": "split", "pages": [8]}                # each named page stands alone
        {"type": "role", "page": 3, "role": <one of ROLES>}
        {"type": "printed_number", "page": 3, "printed_number": 17 | None}
        {"type": "regions", "page": 4, "regions": [{"box": [left, top, right, bottom], "label": ...}, ...]}
        {"type": "revert", "revisions": [2]}           # undo is a further revision; history is never erased

    ``author`` is ``"learner"`` or ``"agent"`` and ``reason`` says why. Each revision names its predecessor by number
    and bytes and lists the earlier corrections it replaces. A logical-slide id that changes stays resolvable as an
    alias; a region id keeps naming its box, and a region that is dropped stays answerable as retired.

    A revert names revisions that are in force. Reverting a revert puts the corrections it had undone back in force,
    and the history says so (``restores``); a revision that is not in force cannot be reverted again.

    The learner's corrections stand: a correction with ``author="agent"`` that would replace, contradict or revert a
    learner's correction in force is refused with ``learner_correction_in_force`` and nothing is recorded — unless
    ``learner_authority=True`` says the agent records it on the learner's explicit say (the record then carries that).
    A learner's revert is such a correction: an agent can neither undo it nor record again what it undid (the same
    group or split, another kind for a join the learner made, any role, number or regions for that page), and what
    the learner's revert put back in force is protected like the learner's own. Where two corrections in force speak
    about the same thing, the one that came into force later decides — a correction put back by undoing its undo
    counts from that moment.

    Raises :class:`SourceMapError`: ``invalid_override`` (nothing is recorded), ``learner_correction_in_force``,
    ``override_conflict`` (another writer took the revision number), ``override_chain_broken`` (the stored revisions
    do not verify) or ``source_unavailable``. Corrections belong to the bytes they were made on: a replaced file
    starts without them.
    """
    if author not in AUTHORS:
        raise SourceMapError("invalid_override", f"author must be one of {', '.join(AUTHORS)}")
    if not isinstance(reason, str) or not reason.strip():
        raise SourceMapError("invalid_override", "a correction states its reason")
    current = load_or_build(settings, course_id, item, lecture=lecture, companions=companions, limits=limits)
    sha = current["source"]["sha256"]
    if sha is None or current["page_count"] is None:
        raise SourceMapError("source_unavailable", "this source has no inspected pages to correct")
    version_dir = _source_dir(settings, course_id, item["source_id"]) / sha[:16]
    records, digests, problems = _load_chain(version_dir / "overrides", sha)
    if problems:
        raise SourceMapError("override_chain_broken", "the stored corrections of this source do not form a verified chain; nothing was added")
    count = current["page_count"]
    ledger, _ = _ledger(records, count)
    force = _in_force(ledger)
    live = [record for record in ledger if record["revision"] in force]
    normal = _validate_op(op, count, [record["revision"] for record in records], force)
    agent_alone = author == "agent" and not learner_authority
    if agent_alone:
        against = _learner_conflicts(normal, ledger, count)
        if against:
            raise SourceMapError("learner_correction_in_force", _refusal(against))
    revision = len(records) + 1
    signature = {"kind": author, "ref": author_ref} | ({"learner_authority": True} if author == "agent" and learner_authority else {})
    record = {"contract": OVERRIDE_CONTRACT, "course_id": course_id, "source": {key: current["source"][key] for key in ("source_id", "version_id", "sha256")}, "revision": revision,
              "supersedes": {"revision": revision - 1, "sha256": digests[-1]} if records else None, "author": signature, "reason": reason.strip()[:1000],
              "created_at": now or datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"), "op": normal, "replaces": _replaced_by(normal, live, count, agent_alone)}
    try:
        publish_exclusive(version_dir / "overrides" / f"r{revision:04d}.json", c.dump_json(record))
    except ExclusiveConflict:
        raise SourceMapError("override_conflict", "another correction was recorded first; read the map again and decide whether this one is still wanted") from None
    return load_or_build(settings, course_id, item, lecture=lecture, companions=companions, limits=limits)


# ---- reading a map -------------------------------------------------------------------------------------------------
def _page_record(source_map: dict, page: Any) -> dict:
    count = source_map.get("page_count") or 0
    if not isinstance(page, int) or isinstance(page, bool) or not 1 <= page <= count or page > len(source_map["pages"]):
        raise SourceMapError("page_out_of_range", f"page {page!r} is not a page of this source (1 to {count})")
    return source_map["pages"][page - 1]


def describe_page(source_map: dict, page: int, *, region: str | None = None) -> dict[str, Any]:
    """What is known about one physical page, and a label that words only that.

    ``PDF page 23 of 91 · build 2 of 3`` always names the physical page. The word for a page of a group says what
    kind of group it is: ``build`` (each page holds all of the page before it and adds to it), ``copy`` (the pages
    draw exactly the same), ``step`` (one slide in several states: something changes between the pages, or the
    drawings were not compared), ``part`` (pages joined by a correction). ``Slide 17`` is added only when the printed
    number is established, ``PPTX slide 12`` only when the alignment matched the page, ``PDF label iv`` only when the
    file defines labels that differ from the index. On an upright reading the printed number is worded
    ``printed page 17``.
    """
    record = _page_record(source_map, page)
    slide = next(s for s in source_map["logical_slides"] if s["logical_slide_id"] == record["logical_slide_id"])
    parts = [f"PDF page {page} of {source_map['page_count']}"]
    if record["build_count"] > 1:
        parts.append(f"{_GROUP_WORD.get(slide['kind'], 'part')} {record['build_index']} of {record['build_count']}")
    if record["printed_number"] is not None:  # a number printed on an upright reading is a page number, not a slide number
        parts.append(f"{'printed page' if (source_map.get('page_layout') or {}).get('kind') == 'document' else 'Slide'} {record['printed_number']}")
    if record["pptx"]:
        parts.append(f"PPTX slide {record['pptx']['ordinal']}")
    if record["pdf_label"] is not None and record["pdf_label"] != str(page):
        parts.append(f"PDF label {record['pdf_label']}")
    chosen = None
    if region is not None:
        chosen = next((r for r in record["regions"] if r["region_id"] == region), None)
        if chosen is None:
            if any(r["region_id"] == region for r in record.get("retired_regions") or []):
                raise SourceMapError("region_retired", f"region {region!r} of page {page} was removed by a correction; resolve() says where it was")
            raise SourceMapError("region_unknown", f"page {page} has no region {region!r}")
        parts.append(f"region {chosen['order']} of {len(record['regions'])}")
    return {"page": page, "page_count": source_map["page_count"], "page_index_base": source_map["page_index_base"], "status": record["status"], "pdf_label": record["pdf_label"],
            "printed_number": record["printed_number"], "pptx": record["pptx"], "logical_slide_id": record["logical_slide_id"], "kind": slide["kind"], "grouping": slide["grouping"],
            "confidence": slide["confidence"], "cumulative": slide.get("cumulative"), "text_visibility": slide.get("text_visibility"),
            "build_index": record["build_index"], "build_count": record["build_count"], "group_pages": slide["pages"], "role": record["role"]["role"],
            "logical_role": slide["role"], "forced_requirements": record["role"]["forced_requirements"], "region": chosen, "regions": len(record["regions"]),
            "source": {key: source_map["source"][key] for key in ("source_id", "version_id", "sha256")}, "map_revision": source_map["map_revision"],
            "parts": parts, "label": " · ".join(parts)}


def resolve(source_map: dict, ref: Any) -> dict[str, Any]:
    """Resolve a logical-slide id, an alias, or a physical page address against a map.

    Accepted: a current id (``current``); any id this source version ever issued (``alias`` — it resolves to the
    slides that now hold its pages); ``page=N`` / ``N`` (``page``, always valid for a page of this version),
    optionally with ``#rK`` for a region (``region``; a region a correction removed answers ``region_retired`` with
    the box it had). An id that is spelled like one of this version but was never issued answers ``unissued``: the
    pages its name spells are given, and nothing is claimed about an earlier grouping. An id of other bytes of the
    source answers ``stale_source_version`` and resolves to nothing.
    """
    answer: dict[str, Any] = {"ref": ref, "status": "unknown", "reason": None, "logical_slide_id": None, "logical_slide_ids": [], "pages": [], "page": None, "region": None,
                              "map_revision": source_map.get("map_revision")}
    by_page = {page["page"]: page for page in source_map["pages"]}
    text = str(ref).strip() if isinstance(ref, (str, int)) and not isinstance(ref, bool) else ""
    match = _PAGE_REF.match(text)
    if match:
        number = int(match.group(1))
        if number not in by_page:
            return {**answer, "reason": "page_out_of_range"}
        record = by_page[number]
        answer.update(status="page", page=number, pages=[number], logical_slide_id=record["logical_slide_id"], logical_slide_ids=[record["logical_slide_id"]])
        if match.group(2):
            region = next((r for r in record["regions"] if r["region_id"] == match.group(2)), None)
            if region is not None:
                answer.update(status="region", region=region)
                return answer
            gone = next((r for r in record.get("retired_regions") or [] if r["region_id"] == match.group(2)), None)
            if gone is not None:  # the page address stays valid; the answer says the region is gone and where it was
                return {**answer, "status": "region_retired", "reason": f"retired by correction r{gone['retired_by_revision']}", "region": gone}
            return {**answer, "status": "unknown", "reason": "region_unknown"}
        return answer
    match = _LOGICAL_ID.match(text)
    if not match:
        return {**answer, "reason": "not_an_address"}
    sha = source_map["source"].get("sha256") or ""
    if match.group(1) != sha[:10]:
        known = next((v for v in source_map.get("other_source_versions") or [] if text in v["logical_slide_ids_not_applicable"]), None)
        return {**answer, "status": "stale_source_version", "reason": "the id belongs to other bytes of this source",
                "stale": {"sha256": known["sha256"] if known else None, "sha256_prefix": match.group(1)}}
    first, last = int(match.group(2)), int(match.group(3) or match.group(2))
    pages = list(range(first, last + 1))
    if first > last or any(n not in by_page for n in pages):
        return {**answer, "reason": "page_out_of_range"}
    targets = list(dict.fromkeys(by_page[n]["logical_slide_id"] for n in pages))
    alias = next((a for a in source_map.get("aliases") or [] if a["alias"] == text), None)
    status = "current" if targets == [text] else "alias" if alias is not None else "unissued"
    answer.update(status=status, logical_slide_id=targets[0], logical_slide_ids=targets, pages=pages, page=pages[0] if len(pages) == 1 else None)
    if status == "alias":
        answer["reason"] = f"retired by correction r{alias['retired_by_revision']}"
    elif status == "unissued":
        answer["reason"] = "this id was never issued for this source version; the slides that hold the pages its name spells are given"
    return answer


def summary(source_map: dict) -> dict[str, Any]:
    """Counts and ids only: raw pages, logical slides and the instructional denominators are separate numbers."""
    pages, slides = source_map["pages"], source_map["logical_slides"]
    page_roles = Counter(page["role"]["role"] for page in pages)
    slide_roles = Counter(slide["role"] for slide in slides)
    unsure_roles = Counter(slide["role"] for slide in slides if slide["role_uncertain"])
    groups = Counter(slide["kind"] for slide in slides if len(slide["pages"]) > 1)
    pptx = source_map.get("pptx")
    inspected = source_map["page_count"] is not None
    return {
        "contract": source_map["contract"], "map_revision": source_map["map_revision"], "source": {key: source_map["source"].get(key) for key in ("source_id", "version_id", "sha256")},
        "inspection": source_map["inspection"]["status"], "inspection_transient": bool(source_map["inspection"].get("transient")),
        "page_layout": (source_map.get("page_layout") or {}).get("kind"),
        # One digest over the text and pictures of all pages: two files with other bytes but the same pages share it.
        "pages_content_sha256": _digest([page.get("content_sha256") for page in pages]) if pages else None,
        "raw_pages": len(pages) if inspected else None, "pages_unknown": sum(1 for page in pages if page["status"] != "ok"),
        "logical_slides": len(slides) if inspected else None,
        "groups": {kind: groups.get(kind, 0) for kind in GROUP_KINDS}, "pages_in_groups": sum(len(slide["pages"]) for slide in slides if len(slide["pages"]) > 1),
        "group_members": [{"logical_slide_id": slide["logical_slide_id"], "kind": slide["kind"], "grouping": slide["grouping"], "confidence": slide["confidence"],
                           "cumulative": slide.get("cumulative"), "pages": slide["pages"], "relations": [entry["relation"] for entry in slide["evidence"]]}
                          for slide in slides if len(slide["pages"]) > 1],
        "uncertain_pairs": sum(1 for pair in source_map.get("pairs") or [] if pair["decision"] == "uncertain"),
        "roles": {"pages": {role: page_roles.get(role, 0) for role in ROLES}, "logical_slides": {role: slide_roles.get(role, 0) for role in ROLES}},
        "roles_uncertain_pages": sum(1 for page in pages if page["role"]["uncertain"]), "roles_uncertain_logical_slides": sum(unsure_roles.values()),
        "denominators": {"raw_pages": len(pages) if inspected else None, "logical_slides": len(slides) if inspected else None,
                         "instructional_logical_slides": sum(slide_roles.get(role, 0) for role in _INSTRUCTIONAL),
                         "instructional_pages": sum(page_roles.get(role, 0) for role in _INSTRUCTIONAL),
                         "objectives_or_summary_logical_slides": slide_roles.get("objectives_or_summary", 0),
                         "unknown_role_logical_slides": slide_roles.get("unknown", 0),
                         "no_forced_requirement_logical_slides": sum(slide_roles.get(role, 0) for role in _NO_FORCED_REQUIREMENT),
                         "no_forced_requirement_uncertain_logical_slides": sum(unsure_roles.get(role, 0) for role in _NO_FORCED_REQUIREMENT)},
        "pdf_labels": source_map["pdf_labels"], "printed_numbers": {"established": source_map["printed_numbers"]["established"],
                                                                    "pages_with_number": sum(1 for page in pages if page["printed_number"] is not None)},
        "pptx": None if pptx is None else {"status": pptx["status"], "slide_count": pptx["slide_count"], "hidden_slides": len(pptx["hidden_slides"]),
                                            "alignment": pptx["alignment"]["status"], "matched_pages": sum(1 for page in pages if page["pptx"])},
        "regions": {"pages_with_regions": sum(1 for page in pages if page["regions"]), "regions": sum(len(page["regions"]) for page in pages)},
        "text_visibility": "not_established",
        "overrides_revision": source_map["overrides"]["revision"], "other_source_versions": len(source_map.get("other_source_versions") or []),
    }
